#!/usr/bin/env python3
"""f25_workbook_reader_test.py - F25: a reviewer had no sanctioned way to read a delivered
workbook and produced a false BLOCKING finding because of it.

THE DEFECT. A blind reviewer reported as BLOCKING that both client workbooks write HTML
entities as literal cell text, naming 29 and 121 affected cells. It did not reproduce: the
reviewer had unzipped the archive and grepped the raw sheet XML, which legitimately carries a
numeric character reference for the unknown-value sentinel. Decoded with a conformant reader
the cell holds exactly one character and Excel renders it correctly (0 of 5,900 string cells
carried an undecoded entity). No rubric said how to read a workbook, so an agent did the
obvious wrong thing. gates.md now states the rule and gives the one-line way to do it; the
two data reviewers' prompts carry the same line.

What this pins:
(1) gates.md has the section, states "decoding reader, never the archive's XML", and carries
    the two one-liners; g-honesty and g-trace carry the read one-liner;
(2) the documented one-liner RUNS: extracted from gates.md verbatim and executed against a
    workbook whose cell holds the sentinel character, it prints the row and no `&#`;
(3) the documented entity test discriminates: 0 on the sentinel workbook (the false finding),
    1 on a workbook whose cell literally holds an entity string (a true finding). The
    sentinel is written by code point on purpose: the house rule forbids authoring the
    character, and here it is a VALUE that must round-trip, not prose.

Run: python evals/f25_workbook_reader_test.py"""
from __future__ import annotations

import os
import re
import shlex
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent

SENTINEL = chr(0x2014)   # the codebase's unknown-value sentinel, as a code point (see docstring)


def _one_liners(gates: str) -> list[str]:
    # every backticked `python -c "..." "<file>.xlsx"` command in the F25 section
    i = gates.find("### Reading a delivered workbook")
    j = gates.find("### What `final_gate.py` enforces", i)
    return re.findall(r"`(python -c \".*?\" \"<file>\.xlsx\")`", gates[i:j], flags=re.S)


def _run(cmd: str, xlsx: Path) -> subprocess.CompletedProcess:
    argv = shlex.split(cmd.replace('"<file>.xlsx"', '"' + str(xlsx).replace("\\", "/") + '"'),
                       posix=True)
    argv[0] = sys.executable
    # force UTF-8 on the child's stdout so the sentinel survives a cp1252 console
    env = dict(os.environ, PYTHONIOENCODING="utf-8")
    return subprocess.run(argv, capture_output=True, text=True, encoding="utf-8",
                          timeout=120, env=env)


def main() -> int:
    fails: list[str] = []

    def check(cond, msg):
        if not cond:
            fails.append(msg)
            print(f"[FAIL] {msg}")
        else:
            print(f"[PASS] {msg}")

    # (1) the docs
    gates = (ROOT / "reference" / "gates.md").read_text(encoding="utf-8")
    check("### Reading a delivered workbook" in gates, "gates.md has the workbook-reading section")
    check("decoding reader, never by unzipping" in gates,
          "gates.md: the rule is 'decoding reader, never the archive'")
    check("NOT literal cell text" in gates,
          "gates.md: says an escape in the XML is not literal cell text")
    check("29 and 121" in gates, "gates.md: records the false finding so the rule has a WHY")
    cmds = _one_liners(gates)
    check(len(cmds) == 2, f"gates.md carries the two one-liners ({len(cmds)} found)")
    for kind in ("g-honesty", "g-trace"):
        p = (ROOT / "prompts" / f"{kind}.md").read_text(encoding="utf-8")
        check("openpyxl.load_workbook" in p and "never by grepping" in p,
              f"{kind}.md carries the decoding-reader line")
    # the section sits BEFORE final_gate's enforcement list, i.e. where a reviewer reads
    check(0 < gates.find("### Reading a delivered workbook") < gates.find("### What `final_gate.py` enforces"),
          "gates.md: the section is in the reviewer-facing part, not buried at the end")

    # (2)+(3) the one-liners run and discriminate
    try:
        from openpyxl import Workbook, load_workbook
    except Exception as e:  # pragma: no cover - the doc itself says what to do without it
        print(f"[SKIP] openpyxl unavailable ({e}); the functional half cannot run here")
        Workbook = None  # type: ignore
    if Workbook is not None and len(cmds) == 2:
        with tempfile.TemporaryDirectory() as td:
            good = Path(td) / "sentinel.xlsx"
            bad = Path(td) / "literal.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Source Ledger"
            ws.append(["field", "value"])
            ws.append(["landPrice", SENTINEL])
            ws.append(["park", "A & B"])
            wb.save(good)
            wb2 = Workbook()
            ws2 = wb2.active
            ws2.append(["field", "value"])
            ws2.append(["landPrice", "&#8212;"])       # literal entity TEXT: a real defect
            wb2.save(bad)

            # the decoded value is one character, whatever the zip's XML looks like
            rb = load_workbook(good, read_only=True)   # read_only holds the zip open: close it,
            v = rb.active["B2"].value                    # or Windows refuses the temp-dir cleanup
            rb.close()
            check(v == SENTINEL and len(v) == 1,
                  "decoded: the sentinel cell holds exactly one character")
            raw = ""
            with zipfile.ZipFile(good) as z:
                for n in z.namelist():
                    if n.startswith("xl/") and n.endswith(".xml"):
                        raw += z.read(n).decode("utf-8", errors="replace")
            check(SENTINEL in raw or "&#8212;" in raw or "&#x2014;" in raw.lower(),
                  "the archive XML carries the sentinel either raw or as a character reference "
                  "(either way it is the XML's business, not the cell's)")

            read_cmd, count_cmd = cmds
            r = _run(read_cmd, good)
            check(r.returncode == 0 and "Source Ledger" in r.stdout and SENTINEL in r.stdout
                  and "&#" not in r.stdout,
                  f"the documented READ one-liner runs and prints the decoded row (rc={r.returncode}"
                  f"{'; ' + r.stderr.strip()[:120] if r.stderr.strip() else ''})")
            r = _run(count_cmd, good)
            check(r.returncode == 0 and r.stdout.strip().startswith("0 cell(s)"),
                  f"the documented ENTITY test reports 0 on the sentinel workbook ({r.stdout.strip()[:60]!r})")
            r = _run(count_cmd, bad)
            check(r.returncode == 0 and r.stdout.strip().startswith("1 cell(s)"),
                  f"...and 1 on a workbook whose cell literally holds an entity ({r.stdout.strip()[:60]!r})")

    print(f"\n{'PASS' if not fails else 'FAIL'} f25_workbook_reader_test ({len(fails)} failure(s))")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
