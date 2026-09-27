#!/usr/bin/env python3
"""master_list_per_file_test.py - master-list brochure rows are per deck FILE, and an answered
per-CLUSTER sheet is carried over without re-asking. (2026-09-26 test run, fix 3.1.)

THE LIVE DEFECT. Deck rows were one per cluster with an id digesting the cluster's file set, so
a label that fused two singleton clusters deleted two answered ids and minted one: the
fingerprint changed, `is_answered` failed and exit 17 re-fired on a finished sheet. The deck
dispatch keyed its skip on the cluster LABEL, so a relabel also made an answered No stop
applying.

WHAT THIS PINS
  1. a two-file cluster yields TWO rows; a singleton's id is byte-identical to the old
     cluster id (no re-ask for the common all-singleton work dir);
  2. regrouping or relabelling the same files leaves expected_hash unchanged;
  3. expected_hash == build_auto's input_hash, including two decks sharing a basename, which
     now get DISTINCT ids (they used to share one);
  4. an old per-cluster master_list.json (a two-file row answered No) migrates: is_answered
     holds for the current inputs, both files are skipped by excluded_deck_files, the old file
     is backed up, `split_from` names the old id;
  5. a sheet answered while a label had fused two decks migrates, run notes re-keyed;
  6. a genuinely new deck file blocks migration (the sheet re-opens) and carry_forward
     pre-fills the known files from master_list.json by file name;
  7. a file under a No row and a Yes row is kept (fail open), in excluded_deck_files and in
     the merge-side _excluded_keys; a headless file pre-fills and excludes nothing;
  8. a PDF and a PPTX sharing a stem are two rows that name each other;
  9. two old rows sharing a basename that DISAGREE are never split by guesswork.
Offline; injected first-page text; no PDFs are opened.
"""
from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "helpers"))
import master_list as ML  # noqa: E402
import master_list_build as MB  # noqa: E402

FAILS: list = []


def ck(ok, msg):
    print(("  [PASS] " if ok else "  [FAIL] ") + msg)
    if not ok:
        FAILS.append(msg)


def _rec(src, park, loc):
    return {"park": park, "city": "Sometown", "postcode": "AB1 2CD", "warehouseArea": 10000,
            "__meta": {"source_file": src, "source_type": "xlsx",
                       "prov": {"park": f"{loc} (tracker)"}}}


RECORDS = {"Tracker.xlsx": [_rec("Tracker.xlsx", "Riverside Park", "Sheet1!B5")]}
A, B, N = "Options - Madrid.pdf", "New stock - Madrid.pdf", "Options - Northtown.pdf"
PAGES = {A: "RIVERSIDE PARK\nMadrid 28001\n12,000 sq m", B: "", N: "NORTH GATE\n"}


def _fpt(p):
    return PAGES.get(Path(str(p)).name, "")


def _work() -> Path:
    return Path(tempfile.mkdtemp(prefix="cbre_mlpf_"))


def _answered(work: Path, rows: list, rec_rows: list, extra: dict | None = None) -> dict:
    """An answered master_list.json in the shape master_list_read writes, hashed the way the
    manifest was hashed when it was built (records + the deck ids ON the sheet)."""
    all_rows = [{"row_id": r["row_id"], "include": "Yes", "run_notes": "", "property": r["property"],
                 "source_type": r["source_type"], "cluster": None,
                 "source_files": r["source_files"], "duplicate_group": ""} for r in rec_rows] + rows
    ml = {"generated": "2026-09-01T10:00:00", "source_workbook": "Master List.xlsx",
          "input_hash": ML.fingerprint(all_rows),
          "counts": {"rows": len(all_rows)}, "rows": all_rows,
          "run_notes": {r["row_id"]: r["run_notes"] for r in all_rows
                        if r.get("include") == "Yes" and r.get("run_notes")},
          "hand_typed": [], "deleted_since_build": []}
    ml.update(extra or {})
    (work / ML.ANSWERS).write_text(json.dumps(ml), encoding="utf-8")
    return ml


def main() -> int:
    rec_rows = ML._record_rows(RECORDS)

    print("== 1. one row per deck FILE ==")
    clusters = {"Madrid": {"pdfs": [A, B], "region": "Madrid"}}
    w = _work()
    auto = ML.build_auto(w, RECORDS, clusters, w, _fpt)
    decks = [r for r in auto["rows"] if r["source_type"] == "Brochure"]
    ck(len(decks) == 2, f"a two-file cluster gives two rows (got {len(decks)})")
    ck(sorted(r["source_files"][0] for r in decks) == sorted([A, B])
       and all(len(r["source_files"]) == 1 for r in decks),
       "each row carries exactly one file")
    ck(all(r["cluster"] == "Madrid" for r in decks), "the cluster label is kept, informational only")
    ck({r["row_id"] for r in decks} == {ML.deck_row_id(A), ML.deck_row_id(B)},
       "row ids are deck_row_id of the file")
    ck(ML.deck_row_id(N) == ML.cluster_row_id("Northtown", [N]) == ML.cluster_row_id(None, [N]),
       "a singleton's per-file id is byte-identical to the old cluster id (no re-ask)")

    print("== 2. regrouping / relabelling the same files changes nothing ==")
    regrouped = {"East": {"pdfs": [A]}, "West": {"pdfs": [B]}}
    ck(ML.expected_hash(RECORDS, clusters) == ML.expected_hash(RECORDS, regrouped)
       == ML.expected_hash(RECORDS, {"Renamed": {"pdfs": [B, A]}}),
       "expected_hash depends on the files, not on how they are grouped or labelled")

    print("== 3. expected_hash mirrors build_auto, basename collisions included ==")
    ck(ML.expected_hash(RECORDS, clusters) == auto["input_hash"],
       "expected_hash == build_auto input_hash for a multi-file cluster")
    coll = {"Site A": {"pdfs": ["site a/photos.pdf"]}, "Site B": {"pdfs": ["site b/photos.pdf"]},
            "Solo": {"pdfs": [N]}}
    w3 = _work()
    auto3 = ML.build_auto(w3, RECORDS, coll, w3, _fpt)
    ids3 = [r["row_id"] for r in auto3["rows"] if r["source_type"] == "Brochure"]
    ck(len(set(ids3)) == 3, f"two decks sharing a basename get DISTINCT ids ({ids3})")
    ck(ML.expected_hash(RECORDS, coll) == auto3["input_hash"],
       "...and expected_hash still equals build_auto's hash")
    ck(ML.deck_row_id(N) in ids3, "a non-colliding file keeps its ordinary id")

    print("== 4. an old per-cluster sheet (two-file row = No) migrates ==")
    w4 = _work()
    old_id = ML.cluster_row_id("Madrid", [A, B])
    _answered(w4, [{"row_id": old_id, "include": "No", "run_notes": "", "property": "Madrid options",
                    "source_type": "Brochure", "cluster": "Madrid", "source_files": [A, B],
                    "duplicate_group": ""}], rec_rows)
    exp4 = ML.expected_hash(RECORDS, clusters)
    ck(not ML.is_answered(w4, exp4), "before migration the per-file hash does not match")
    ck(ML.migrate_per_file(w4, RECORDS, clusters) is True, "migrate_per_file re-keys the sheet")
    ck(ML.is_answered(w4, exp4), "...after which is_answered holds for the current inputs")
    ck(ML.excluded_deck_files(w4) == {A.lower(), B.lower()},
       f"both files are skipped by the deck dispatch ({sorted(ML.excluded_deck_files(w4))})")
    ml4 = ML.load_answers(w4)
    split = [r for r in ml4["rows"] if str(r.get("row_id")).startswith("deck:")]
    ck(len(split) == 2 and all(r.get("split_from") == old_id for r in split),
       "two per-file rows, each with split_from = the old id")
    ck((w4 / ML.PRE_SPLIT_BACKUP).exists()
       and json.loads((w4 / ML.PRE_SPLIT_BACKUP).read_text(encoding="utf-8"))["input_hash"]
       != ml4["input_hash"], "the old file is backed up before the rewrite")
    ck((ml4.get("migrated") or {}).get("rows_split") == 1
       and ml4["counts"]["excluded"] == 2 and ml4["counts"]["included"] == 1,
       f"migrated stamp + recomputed counts ({ml4.get('migrated')}, {ml4['counts']})")
    ck(ML.migrate_per_file(w4, RECORDS, clusters) is False,
       "a second call is a no-op (nothing left to re-key)")

    print("== 5. a sheet answered while a label had FUSED two decks migrates ==")
    w5 = _work()
    fused_id = ML.cluster_row_id("Northtown", [A, N])
    _answered(w5, [{"row_id": fused_id, "include": "Yes", "run_notes": "use the mezzanine figure",
                    "property": "Northtown", "source_type": "Brochure", "cluster": "Northtown",
                    "source_files": [A, N], "duplicate_group": ""}], rec_rows)
    now5 = {"Madrid": {"pdfs": [A]}, "Northtown": {"pdfs": [N]}}
    ck(ML.migrate_per_file(w5, RECORDS, now5) is True, "the fused sheet migrates")
    ml5 = ML.load_answers(w5)
    ck(ML.is_answered(w5, ML.expected_hash(RECORDS, now5)), "...and is answered for the split inputs")
    ck(ml5["run_notes"].get(ML.deck_row_id(A)) == "use the mezzanine figure"
       and ml5["run_notes"].get(ML.deck_row_id(N)) == "use the mezzanine figure"
       and fused_id not in ml5["run_notes"],
       "the run note is re-keyed to every derived id")

    print("== 6. a new deck file blocks migration; carry_forward pre-fills by file name ==")
    w6 = _work()
    _answered(w6, [{"row_id": old_id, "include": "No", "run_notes": "", "property": "Madrid",
                    "source_type": "Brochure", "cluster": "Madrid", "source_files": [A, B],
                    "duplicate_group": ""}], rec_rows)
    plus = {"Madrid": {"pdfs": [A, B]}, "Northtown": {"pdfs": [N]}}
    ck(ML.migrate_per_file(w6, RECORDS, plus) is False,
       "a new file means the sheet answered other inputs: no migration")
    ck(not ML.is_answered(w6, ML.expected_hash(RECORDS, plus)), "...so exit 17 re-opens")
    auto6 = ML.build_auto(w6, RECORDS, plus, w6, _fpt)
    rows6 = [dict(r) for r in auto6["rows"]]
    for r in rows6:
        r["include"], r["run_notes"] = "", ""
    kept, _manual = MB.carry_forward(w6 / "no such workbook.xlsx", rows6, w6)
    got = {r["source_files"][0]: r["include"] for r in rows6 if r["source_type"] == "Brochure"}
    ck(kept == 2 and got.get(A) == "No" and got.get(B) == "No" and got.get(N) == "",
       f"the two known files are pre-filled from master_list.json, the new one is blank ({got})")
    MB.build(w6)
    from openpyxl import load_workbook  # noqa: E402
    ws = load_workbook(w6 / ML.WORKBOOK)[MB.SHEET]
    inc = {str(ws.cell(r, MB.IDX["row_id"]).value): ws.cell(r, MB.IDX["include"]).value
           for r in range(MB.FIRST_ROW, ws.max_row + 1) if ws.cell(r, MB.IDX["row_id"]).value}
    ck(inc.get(ML.deck_row_id(A)) == "No" and not inc.get(ML.deck_row_id(N)),
       "build() passes the work dir, so the workbook shows the carried answers")

    print("== 7. No + Yes on one file keeps it; headless is nobody's answer ==")
    w7 = _work()
    (w7 / ML.ANSWERS).write_text(json.dumps({"input_hash": "x", "rows": [
        {"row_id": old_id, "include": "No", "source_files": [A, B]},
        {"row_id": ML.deck_row_id(A), "include": "Yes", "source_files": [A]}]}), encoding="utf-8")
    ck(ML.excluded_deck_files(w7) == {B.lower()},
       f"a file under a Yes row is read even if a legacy No row covers it ({ML.excluded_deck_files(w7)})")
    _ids, _files = ML._excluded_keys(ML.load_answers(w7))
    ck(A.lower() not in _files and B.lower() in _files, "the merge-side filter agrees")
    w7h = _work()
    ML.write_headless(w7h, auto, "headless (eval)")
    ck(ML.excluded_deck_files(w7h) == set() and MB._answers_by_deck_file(w7h) == {},
       "a headless file excludes nothing and pre-fills nothing")

    print("== 8. a PDF and a PPTX of one deck are two rows that name each other ==")
    w8 = _work()
    twin = {"Riverside": {"pdfs": ["Riverside Park.pdf"], "pptxs": ["Riverside Park.pptx"]}}
    rows8 = [r for r in ML.build_auto(w8, {}, twin, w8, _fpt)["rows"]]
    ck(len(rows8) == 2, "two rows")
    ck(all("also on disk as" in r["notes"] and "answer both the same way" in r["notes"]
           for r in rows8), "each notes the other copy")
    ck(all(not r.get("duplicate_group") for r in rows8),
       "...and no automatic duplicate group (it would seed 'same' across their records)")

    print("== 9. old rows that DISAGREE about one basename are not split by guesswork ==")
    w9 = _work()
    p1, p2 = "site a/photos.pdf", "site b/photos.pdf"
    oid = ML.cluster_row_id(None, ["photos.pdf"])
    _answered(w9, [{"row_id": oid, "include": "Yes", "run_notes": "", "property": "A",
                    "source_type": "Brochure", "cluster": "Site A", "source_files": ["photos.pdf"],
                    "duplicate_group": ""},
                   {"row_id": oid, "include": "No", "run_notes": "", "property": "B",
                    "source_type": "Brochure", "cluster": "Site B", "source_files": ["photos.pdf"],
                    "duplicate_group": ""}], rec_rows)
    ck(ML.migrate_per_file(w9, RECORDS, {"Site A": {"pdfs": [p1]}, "Site B": {"pdfs": [p2]}})
       is False, "a Yes and a No under one shared id cannot be assigned: the sheet re-opens")

    print()
    if FAILS:
        print(f"MASTER LIST PER-FILE TEST: FAIL ({len(FAILS)})")
        return 1
    print("MASTER LIST PER-FILE TEST: PASS")
    return 0


if __name__ == "__main__":
    sys.exit(main())
