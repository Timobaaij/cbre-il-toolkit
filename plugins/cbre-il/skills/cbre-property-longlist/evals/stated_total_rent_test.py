#!/usr/bin/env python3
"""stated_total_rent_test.py - v46: a stated annual TOTAL rent has a canonical home.

A brochure that quotes "£750,000 per annum exclusive" and no per-area rate used to have nowhere to
put it: it shipped only as an extras column, and the Total annual rent row read tbd. v46 adds the
optional string field `quotingRentTotal` and falls back to it - verbatim, marked as stated - in the
Total annual rent outputs ONLY when no GLA x rate total can be computed:
  * the Longlist workbook (deliver._total_rent / value_for("__total_annual")) -> "... (as stated)";
  * the dashboard (totalRentStr -> quotedTotalRent), with no basis line under it.
Nothing is derived from it: no monthly figure, no per-area rate, warehouseRent stays tbd. With a
rate present the computed GLA x rate total is unchanged and wins over the stated one.

Offline. Needs node to execute the chrome (as f32_rent_basis_test does).
Usage: python evals/stated_total_rent_test.py
"""
from __future__ import annotations

import csv
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "helpers"))
import _common as C  # noqa: E402
import build_dashboard as BD  # noqa: E402
import deliver as D  # noqa: E402

PX = ("data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0"
      "lEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")
STATED = "\u00a3750,000 per annum exclusive"
FAILS: list[str] = []


def ck(ok, msg):
    print(("  [PASS] " if ok else "  [FAIL] ") + msg)
    if not ok:
        FAILS.append(msg)


def _props(all_stated: bool = False):
    base = {"country": "GB", "developer": "Dev", "city": "Sometown", "status": "Available",
            "photo": PX, "gallery": [PX], "lat": 52.49, "lng": -0.69, "areaUnit": "sq ft"}
    stated = dict(base, id=1, park="StatedOnly", warehouseArea=40000, warehouseRent=C.BLANK,
                  quotingRentTotal=STATED)
    if all_stated:
        return [stated, dict(stated, id=2, park="StatedToo")]
    rated = dict(base, id=2, park="Rated", warehouseArea=40000, rentUnit="\u00a3/sq ft/yr",
                 warehouseRent="\u00a39.50 / sq ft / year", warehouseRentVal=9.5)
    both = dict(rated, id=3, park="RatedAndStated", quotingRentTotal=STATED)
    return [stated, rated, both]


def _canon(all_stated: bool = False):
    return {"meta": {"client": "StatedTotal",
                     "units": {"area": "sq ft", "rent": "\u00a3/sq ft/yr"},
                     "hero": {"topbar_meta": "", "eyebrow": "", "title_html": "",
                              "lede": "", "footer_copyright": ""}},
            "pois": [], "regions": {}, "properties": _props(all_stated)}


JS = r"""
import fs from 'node:fs';
import vm from 'node:vm';
const html = fs.readFileSync(process.argv[2], 'utf8');
const STATED = process.argv[3];
const scripts = [...html.matchAll(/<script(?![^>]*\bsrc=)[^>]*>([\s\S]*?)<\/script>/gi)].map(m => m[1]);
const NAMES = ['PROPS', 'BLANK', 'totalRentStr', 'totalRentHTML', 'detailHTML', 'compareHTML', 'T'];
const code = scripts.join('\n;\n') + '\n;\n' +
  NAMES.map(n => `__capture__(${JSON.stringify(n)}, typeof ${n} !== "undefined" ? ${n} : undefined);`).join('\n') + '\n';
const sink = new Proxy(function () {}, {
  get: (_t, p) => (p === Symbol.toPrimitive || p === 'toString' || p === 'valueOf') ? () => '' : sink,
  apply: () => sink, construct: () => sink, has: () => true,
});
const target = { console };
for (const name of Object.getOwnPropertyNames(globalThis)) {
  if (!(name in target)) { try { target[name] = globalThis[name]; } catch { /* ignore */ } }
}
target.globalThis = target;
target.__capture__ = (name, val) => { target[name] = val; };
const ctx = vm.createContext(new Proxy(target, { get: (t, p) => (p in t ? t[p] : sink), has: () => true }));
try { vm.runInContext(code, ctx, { filename: 'built.inline.js' }); }
catch (e) { console.error('FAIL: template script threw during eval:', e && e.message); process.exit(1); }
const { PROPS: props, totalRentStr, totalRentHTML, detailHTML, compareHTML, T } = target;
const fails = [];
const ck = (ok, label) => { console.log((ok ? '  ok   ' : '  FAIL ') + label); if (!ok) fails.push(label); };
const byPark = n => props.find(p => p.park === n);
const s = byPark('StatedOnly');
if (process.argv[4] === 'all') {
  const cmp = compareHTML(props);
  ck(cmp.includes(T('row_total_annual_rent')),
     'Compare keeps the Total annual rent row when only stated totals exist (FIELD_PRESENT keys on quotingRentTotal too)');
  ck(cmp.includes(STATED), 'Compare shows the stated total in that row');
} else {
  const r = byPark('Rated'), b = byPark('RatedAndStated');
  ck(totalRentStr(s, false) === STATED, 'no rate: Total annual rent is the stated total VERBATIM');
  ck(totalRentStr(s, true) === null, 'no rate: NO monthly figure is derived from a stated string');
  ck(!String(totalRentHTML(s, false)).includes('rent-basis'), 'no rate: the stated total carries no area x rate basis line');
  ck(detailHTML(s).includes(STATED), 'the modal prints the stated total under Total annual rent');
  ck(totalRentStr(r, false) === '\u00a3 380,000 / yr', 'rate present: the computed GLA x rate total is unchanged');
  ck(totalRentStr(b, false) === '\u00a3 380,000 / yr' && !String(totalRentHTML(b, false)).includes(STATED),
     'rate AND stated total: the computed total wins, the stated one is not substituted');
  ck(String(totalRentHTML(b, false)).includes('rent-basis'), 'rate present: the basis line still prints');
}
process.exit(fails.length ? 1 : 0);
"""


def _run_chrome(node: str, hp: Path, mode: str, td: Path) -> None:
    mjs = td / "stated_total_rent.mjs"
    mjs.write_text(JS, encoding="utf-8")
    r = subprocess.run([node, str(mjs), str(hp), STATED, mode], capture_output=True, text=True,
                       encoding="utf-8", errors="replace")
    print((r.stdout or "").rstrip())
    if r.returncode != 0:
        print((r.stderr or "").rstrip())
    ck(r.returncode == 0, f"the EXECUTED chrome assertions pass ({mode})")


def main() -> int:
    print("== the field is registered like warehouseRent / serviceCharge ==")
    schema = json.loads((ROOT / "templates" / "canonical.schema.json").read_text(encoding="utf-8"))
    node_def = schema["$defs"]["property"]["properties"].get("quotingRentTotal") or {}
    ck(node_def.get("type") == "string", "canonical.schema.json declares quotingRentTotal as a string")
    ck("warehouseRent" in node_def.get("x-reader-format", ""),
       "its reader format says it never goes into warehouseRent")
    ck("quotingRentTotal" in C.STRING_FIELDS, "_common.STRING_FIELDS carries it (coerced, sentinel-filled, ledger-traced)")
    ck("quotingRentTotal" in C.IDENTIFIER_FIELDS and not C.is_translatable_value("quotingRentTotal", STATED),
       "it is an IDENTIFIER field: a figure+currency string is never sent to the translator")
    ck("quotingRentTotal" in C.canonical_property_fields(), "it is a canonical property field (reader-fillable)")

    print()
    print("== merge keeps it verbatim and derives no rate from it ==")
    try:
        import merge as M
        rec = M.canonicalize({"id": 1, "park": "StatedOnly", "country": "GB", "areaUnit": "sq ft",
                              "warehouseArea": 40000, "quotingRentTotal": STATED})
        C.fill_render_sentinels(rec)
        ck(rec.get("quotingRentTotal") == STATED, "quotingRentTotal survives canonicalize VERBATIM")
        ck(C._N.looks_unknown(rec.get("warehouseRent")), f"warehouseRent stays tbd ({rec.get('warehouseRent')!r})")
        ck(not isinstance(rec.get("warehouseRentVal"), (int, float)),
           f"no per-area rate is back-derived ({rec.get('warehouseRentVal')!r})")
    except Exception as e:  # merge is a separate module; report, never mask
        ck(False, f"merge.canonicalize ran ({type(e).__name__}: {e})")

    print()
    print("== the Longlist workbook's Total annual rent ==")
    props = {p["park"]: p for p in _props()}
    ck(D._total_rent(props["StatedOnly"], False) == f"{STATED} (as stated)",
       "no rate: Total annual rent = the stated total, marked (as stated)")
    ck(D._total_rent(props["StatedOnly"], True) == C.BLANK, "no rate: Total monthly rent stays tbd")
    ck(D._total_rent(props["Rated"], False) == "\u00a3 380,000 / yr",
       "rate present: the derived GLA x rate total is unchanged")
    ck(D._total_rent(props["RatedAndStated"], False) == "\u00a3 380,000 / yr",
       "rate AND stated total: the computed total wins")
    with tempfile.TemporaryDirectory() as td:
        out = Path(td) / "Longlist.xlsx"
        D.longlist_xlsx(_canon(), out)
        if out.exists():
            from openpyxl import load_workbook
            ws = load_workbook(out).active
            rows = [[c.value for c in r] for r in ws.iter_rows()]
        else:
            with open(out.with_suffix(".csv"), encoding="utf-8") as fh:
                rows = list(csv.reader(fh))
        hdr = rows[0]
        rec = {r[hdr.index("Property / Park")]: r for r in rows[1:]}
        s_row = rec["StatedOnly"]
        ck(s_row[hdr.index("Total annual rent")] == f"{STATED} (as stated)",
           "the Longlist's Total annual rent cell shows the stated total")
        ck(s_row[hdr.index("Warehouse rent (annual)")] == C.BLANK,
           "the Longlist's warehouse rent stays tbd (no rate invented)")
        ck(rec["Rated"][hdr.index("Total annual rent")] == "\u00a3 380,000 / yr",
           "the Longlist's computed total is unchanged for a rated property")

    print()
    print("== the dashboard chrome, executed ==")
    node = shutil.which("node") or r"C:\Users\TBaaij\nodejs\node.exe"
    if not Path(node).exists() and not shutil.which("node"):
        ck(False, "node is required to execute the chrome (install node or add it to PATH)")
    else:
        for mode in ("mixed", "all"):
            with tempfile.TemporaryDirectory() as td:
                d = Path(td)
                cp, hp = d / "c.json", d / "b.html"
                cp.write_text(json.dumps(_canon(mode == "all")), encoding="utf-8")
                BD.build(cp, hp)
                _run_chrome(node, hp, mode, d)

    print()
    if FAILS:
        print(f"STATED TOTAL RENT TEST: FAIL ({len(FAILS)})")
        return 1
    print("STATED TOTAL RENT TEST: PASS")
    return 0


if __name__ == "__main__":
    sys.exit(main())
