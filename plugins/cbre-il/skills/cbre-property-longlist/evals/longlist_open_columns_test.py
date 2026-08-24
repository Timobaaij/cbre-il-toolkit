#!/usr/bin/env python3
"""longlist_open_columns_test.py - the flat Longlist workbook is the COMPLETE view.

Open-captured scalar fields (crossDock, buildType, drive-time columns, ...) reach
canonical.json but had no column in the fixed LONGLIST_COLUMNS list, so the one
artefact billed as 'one property per ROW, variables in COLUMNS' silently omitted
them. Pinned here: every non-denied scalar field on any property ships as an extra
right-hand column with a prettified header; media/derived/internal keys stay out;
the fixed column order is untouched. Fast + offline (needs openpyxl).
"""
import sys
import pathlib
import tempfile

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent / "helpers"))
import deliver  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


def check(name, cond):
    if not cond:
        raise AssertionError(name)


canonical = {
    "meta": {"units": {"rent": "£/sq ft/yr"}},
    "properties": [
        {"id": "p1", "park": "Alpha Park", "city": "Northampton",
         "warehouseArea": 61262, "areaUnit": "sq m",
         "buildType": "Second hand", "crossDock": "Yes",
         "truckMilesToNationalHub": "32.2",
         "photo": "data:image/jpeg;base64,xxxx", "preBaked": True,
         "warehouseRentVal": 9.75,
         "warehouseRent": "£9.75 / sq ft / year"},
        {"id": "p2", "park": "Beta Park", "city": "Rugby",
         "warehouseArea": 30000, "areaUnit": "sq m"},
    ],
}

out = pathlib.Path(tempfile.mkdtemp()) / "Longlist.xlsx"
deliver.longlist_xlsx(canonical, out)
ws = load_workbook(str(out))["Longlist"]
headers = [c.value for c in ws[1]]

# fixed order untouched at the front
fixed = [h for _, h in deliver.LONGLIST_COLUMNS]
check("fixed-prefix", headers[:len(fixed)] == fixed)

# open scalars appended, prettified
check("col-buildtype", "Build type" in headers)
check("col-crossdock", "Cross dock" in headers)
check("col-truckmiles", "Truck miles to national hub" in headers)

# denied keys stay out
check("no-photo", not any(h and "photo" in str(h).lower() for h in headers))
check("no-prebaked", not any(h and "baked" in str(h).lower() for h in headers))
# merge stamps warehouseRent on every property; it must NOT duplicate the fixed
# "Warehouse rent (annual)" column as a bare "Warehouse rent" extra
check("no-dup-rent", "Warehouse rent" not in headers)

# values land in the right columns; a property lacking the field reads 'tbd'
rows = {r[headers.index("Property / Park")]: r
        for r in ([c.value for c in row] for row in ws.iter_rows(min_row=2))}
check("v-crossdock", rows["Alpha Park"][headers.index("Cross dock")] == "Yes")
check("v-buildtype", rows["Alpha Park"][headers.index("Build type")] == "Second hand")
check("v-tbd", rows["Beta Park"][headers.index("Cross dock")] == "tbd")

print("LONGLIST OPEN COLUMNS TEST: PASS")
