#!/usr/bin/env python3
"""Source Ledger for the CBRE I&L Occupier Brief.

Every material claim in the brief has a row here. The ledger is the audit trail
that makes the brief defensible three weeks after the meeting, when somebody asks
where a number came from.

Usage:
    python ledger.py init      <ledger.csv>
    python ledger.py validate  <ledger.csv>
    python ledger.py merge     <out.csv> <in1.csv> [in2.csv ...]
    python ledger.py xlsx      <ledger.csv> [out.xlsx]

Pure standard library except `xlsx`, which uses openpyxl if it is installed and
otherwise tells you the CSV is the deliverable.
"""
from __future__ import annotations

import csv
import os
import re
import sys
from datetime import date

COLUMNS = [
    "claim_id",           # C001, C002 ... unique
    "section",            # the brief section the claim lands in
    "claim",              # the clause AS IT APPEARS in the brief, verbatim
    "figure_at_source",   # the number or quote exactly as the source states it
    "source_title",
    "publisher",
    "source_url",
    "tier",               # 1 to 6, see reference/evidence-and-ledger.md
    "publication_date",   # YYYY-MM-DD, or YYYY-MM if that is all the source gives
    "retrieved_date",     # YYYY-MM-DD
    "sourcing",           # public | proprietary
    "confidence",         # high | medium | low
    "verify_before_use",  # yes | no
    "notes",
]

REQUIRED = [
    "claim_id", "section", "claim", "figure_at_source", "source_title",
    "source_url", "tier", "publication_date", "retrieved_date", "sourcing",
    "confidence", "verify_before_use",
]

TIERS = {"1", "2", "3", "4", "5", "6"}
SOURCING = {"public", "proprietary"}
CONFIDENCE = {"high", "medium", "low"}
YESNO = {"yes", "no"}
DATE_RE = re.compile(r"^\d{4}(-\d{2}(-\d{2})?)?$")
URL_RE = re.compile(r"^https?://\S+$", re.I)

# A financial or footprint figure cannot rest on a single low-tier source.
HARD_FIGURE_SECTIONS = ("at a glance",)


def longpath(path):
    if os.name != "nt":
        return path
    absolute = os.path.abspath(path)
    if len(absolute) < 240 or absolute.startswith("\\\\?\\"):
        return absolute
    return "\\\\?\\" + absolute.replace("/", "\\")


def read(path):
    with open(longpath(path), newline="", encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))


def write(path, rows):
    d = os.path.dirname(os.path.abspath(path))
    if d and not os.path.isdir(longpath(d)):
        os.makedirs(longpath(d), exist_ok=True)
    with open(longpath(path), "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=COLUMNS, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({c: (r.get(c) or "") for c in COLUMNS})


def cmd_init(path):
    example = {
        "claim_id": "C001",
        "section": "At a glance",
        "claim": "Revenue FY25 GBP 4.2bn, up 6.1 per cent",
        "figure_at_source": "Revenue 4,214.3 GBPm (FY24: 3,972.1)",
        "source_title": "Annual Report and Accounts 2025, page 12",
        "publisher": "Testco plc",
        "source_url": "https://example.com/ar2025.pdf",
        "tier": "1",
        "publication_date": "2026-05-14",
        "retrieved_date": date.today().isoformat(),
        "sourcing": "public",
        "confidence": "high",
        "verify_before_use": "no",
        "notes": "Delete this example row.",
    }
    write(path, [example])
    print("Wrote ledger scaffold with 1 example row: %s" % path)
    return 0


def validate(rows, strict_ids=True):
    problems = []
    seen = set()
    for i, r in enumerate(rows, start=2):
        cid = (r.get("claim_id") or "").strip()
        where = "row %d (%s)" % (i, cid or "no claim_id")
        for col in REQUIRED:
            if not (r.get(col) or "").strip():
                problems.append("%s: missing %s" % (where, col))
        if strict_ids:
            if cid and cid in seen:
                problems.append("%s: duplicate claim_id" % where)
            seen.add(cid)
            if cid and not re.match(r"^C\d{3,}$", cid):
                problems.append("%s: claim_id must look like C001" % where)
        tier = (r.get("tier") or "").strip()
        if tier and tier not in TIERS:
            problems.append("%s: tier '%s' is not 1 to 6" % (where, tier))
        for col, allowed in (("sourcing", SOURCING), ("confidence", CONFIDENCE),
                             ("verify_before_use", YESNO)):
            v = (r.get(col) or "").strip().lower()
            if v and v not in allowed:
                problems.append("%s: %s '%s' must be one of %s"
                                % (where, col, v, "/".join(sorted(allowed))))
        for col in ("publication_date", "retrieved_date"):
            v = (r.get(col) or "").strip()
            if v and not DATE_RE.match(v):
                problems.append("%s: %s '%s' must be YYYY-MM-DD" % (where, col, v))
        url = (r.get("source_url") or "").strip()
        if url and not URL_RE.match(url):
            problems.append("%s: source_url is not a retrievable http(s) URL" % where)
        sec = (r.get("section") or "").strip().lower()
        if any(k in sec for k in HARD_FIGURE_SECTIONS) and tier in ("4", "5", "6"):
            if "corroborat" not in (r.get("notes") or "").lower():
                problems.append(
                    "%s: an At a glance figure on a tier-%s source needs a second "
                    "corroborating source named in notes" % (where, tier))
        if (r.get("confidence") or "").strip().lower() == "low" \
                and (r.get("verify_before_use") or "").strip().lower() != "yes":
            problems.append("%s: low confidence must be verify_before_use=yes" % where)
    return problems


def cmd_validate(path):
    rows = read(path)
    problems = validate(rows)
    print("Rows: %d" % len(rows))
    tiers = {}
    for r in rows:
        tiers[r.get("tier", "?")] = tiers.get(r.get("tier", "?"), 0) + 1
    print("Tier mix: " + ", ".join("tier %s: %d" % (k, tiers[k]) for k in sorted(tiers)))
    primary = sum(1 for r in rows if (r.get("tier") or "") in ("1", "2", "3"))
    share = (100.0 * primary / len(rows)) if rows else 0.0
    print("Primary (tier 1 to 3): %d of %d (%.0f per cent)" % (primary, len(rows), share))
    if len(rows) < 15:
        problems.append("ledger holds %d rows; a full brief traces at least 15 distinct sources"
                        % len(rows))
    if rows and share < 50:
        problems.append("only %.0f per cent of rows are tier 1 to 3; the brief is leaning on "
                        "secondary reporting" % share)
    for p in problems:
        print("FAIL  " + p)
    print("STATUS: " + ("PASS" if not problems else "FAIL (%d)" % len(problems)))
    return 0 if not problems else 1


def cmd_merge(out, inputs):
    rows, seen = [], set()
    for path in inputs:
        for r in read(path):
            key = ((r.get("source_url") or "").strip().lower(),
                   (r.get("claim") or "").strip().lower())
            if key in seen:
                continue
            seen.add(key)
            rows.append(r)
    for n, r in enumerate(rows, start=1):
        r["claim_id"] = "C%03d" % n
    write(out, rows)
    print("Merged %d file(s) into %s: %d rows after de-duplication"
          % (len(inputs), out, len(rows)))
    return 0


def cmd_xlsx(path, out=None):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        print("openpyxl is not installed; ship the CSV, it is a valid deliverable.",
              file=sys.stderr)
        return 3
    rows = read(path)
    out = out or os.path.splitext(path)[0] + ".xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Source Ledger"
    ws.append([c.replace("_", " ").title() for c in COLUMNS])
    for r in rows:
        ws.append([(r.get(c) or "") for c in COLUMNS])
    head = Font(bold=True, color="FFFFFF", name="Calibri")
    fill = PatternFill("solid", fgColor="006A4D")
    for cell in ws[1]:
        cell.font = head
        cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    widths = {"claim": 52, "figure_at_source": 46, "source_title": 38,
              "source_url": 44, "notes": 34, "section": 26}
    for i, c in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = widths.get(c, 16)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(longpath(out))
    print("Wrote %s (%d rows)" % (out, len(rows)))
    return 0


USAGE = __doc__


def main(argv):
    if len(argv) < 3:
        print(USAGE)
        return 1
    cmd, rest = argv[1], argv[2:]
    if cmd == "init":
        return cmd_init(rest[0])
    if cmd == "validate":
        return cmd_validate(rest[0])
    if cmd == "merge":
        return cmd_merge(rest[0], rest[1:])
    if cmd == "xlsx":
        return cmd_xlsx(rest[0], rest[1] if len(rest) > 1 else None)
    print(USAGE)
    return 1


if __name__ == "__main__":
    sys.exit(main(sys.argv))
