"""Dump the user's Expense Type and Attendees columns out of Expenses.xlsx.

Usage:
    python read_types.py "<output_dir>/Expenses.xlsx" [claim.json]

Prints one row per claim line: number, date, merchant, claim amount, the
expense CODE to type into EXPENSE_TYPE$N, and the attendee cell verbatim. Pass
claim.json as a second argument to fold both columns back into it, so a rebuild
does not lose them.

This script does not interpret the attendee text. It is free text a person
wrote, and turning "John Doe (Hillwood)" into the "Surname,Firstname" plus
company that PeopleSoft wants is a reading job, not a regex job: names carry
middle initials, particles, double-barrels and inconsistent separators, and a
parser that guesses wrong files a real person's name incorrectly. Read the cell
and decide.

What it does check, because both are silent filing errors:
  - a blank or unrecognised Expense Type
  - an entertaining line with an empty Attendees cell, which PeopleSoft will
    refuse to save

Either makes the exit status non-zero.
"""

import io
import json
import sys

import openpyxl

TYPE_COL, ATTENDEE_COL = 11, 12

VALID_CODES = {
    "SUBSIST", "CLENT", "STFENT", "WRKLNCH", "PUBTRN", "TAXI", "PARKING",
    "TOLLCNG", "RENTCAR", "CARFUEL", "HOTEL", "AIRFARE", "AIRFDOM", "PHONECM",
    "POSTAGE", "CONFSEM", "MBRSHIP", "ITEQUIP", "LEGALPR", "OTHER",
}
NEEDS_ATTENDEE = ("CLENT", "STFENT", "WRKLNCH")


def main():
    if len(sys.argv) not in (2, 3):
        print(__doc__)
        return 1
    xlsx = sys.argv[1]
    ws = openpyxl.load_workbook(xlsx, data_only=True)["Expenses"]

    rows, errors = [], []
    for r in range(2, ws.max_row + 1):
        no = ws.cell(r, 1).value
        if not isinstance(no, int):
            break                      # blank gap, then the totals block
        raw = (ws.cell(r, TYPE_COL).value or "").strip()
        code = raw.split(" - ")[0].strip().upper() if raw else ""
        att = (ws.cell(r, ATTENDEE_COL).value or "").strip()

        if not code:
            errors.append("line %s: Expense Type is blank" % no)
        elif code not in VALID_CODES:
            errors.append("line %s: %r is not an expense type code" % (no, code))
        if code in NEEDS_ATTENDEE and not att:
            errors.append("line %s: %s will not save without an attendee" % (no, code))

        d = ws.cell(r, 2).value
        rows.append({
            "no": no,
            "date": d.strftime("%Y-%m-%d") if hasattr(d, "strftime") else d,
            "merchant": ws.cell(r, 3).value,
            "description": ws.cell(r, 4).value,
            "amount": ws.cell(r, 8).value,
            "ccy": ws.cell(r, 9).value,
            "type_raw": raw,
            "code": code,
            "attendees": att,
        })

    w = max([len(str(x["merchant"] or "")) for x in rows] + [8])
    for x in rows:
        print("%2s  %s  %-*s  %8.2f %s  %-8s %s" % (
            x["no"], x["date"], w, x["merchant"], x["amount"] or 0,
            x["ccy"], x["code"] or "?", x["attendees"]))

    counts = {}
    for x in rows:
        counts[x["code"]] = counts.get(x["code"], 0) + 1
    print("\n%d lines: %s" % (
        len(rows), ", ".join("%s x%d" % (k or "?", v) for k, v in sorted(counts.items()))))

    if len(sys.argv) == 3:
        path = sys.argv[2]
        claim = json.load(io.open(path, encoding="utf-8"))
        by_no = {x["no"]: x for x in rows}
        # claim.json is in the same date-stable order the builder numbered.
        for i, line in enumerate(sorted(claim["lines"], key=lambda l: l["date"]), start=1):
            if i in by_no:
                line["expense_type"] = by_no[i]["type_raw"]
                line["attendees"] = by_no[i]["attendees"]
        io.open(path, "w", encoding="utf-8").write(json.dumps(claim, indent=2))
        print("folded back into %s" % path)

    for e in errors:
        print("ERROR %s" % e, file=sys.stderr)
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
