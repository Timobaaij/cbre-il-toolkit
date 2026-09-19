#!/usr/bin/env python3
"""
Stage 2.5a - build the MASTER LIST: one inventory row per property named anywhere in this
run, for the user to accept or reject before the run proper starts.

WHY THIS STEP EXISTS
--------------------
A requirement's options arrive by three routes at once and the routes overlap badly:

  * the Kato longlist, which surfaces one match request PER BROKER, so the identical
    physical unit arrives two or three times under different agencies;
  * the broker email export, where an agent who has already posted a unit to Kato mails the
    same unit again, often alongside three more that are on Kato under someone else;
  * whatever extra files the user drops in the working directory, which are usually the
    brochures for options already in one of the first two.

Nothing upstream reconciles those. Before this step the run simply took every Kato match,
silently auto-merged the ones that happened to share a postal code and an exact floor area
(common.dedupe_props), and never saw an option that existed only in an email. The user found
out what the run had decided by reading the finished dashboard.

So: enumerate everything first, put the overlaps in front of the user, and let them say yes
or no per row with a free-text instruction to the run. That sheet is then the single source
of truth for what the run builds. Read back by master_list_read.py.

Yes or no, with no third option. A deferred answer has to be resolved before the run can start
anyway, so carrying it in the sheet only moves the same decision further down the pipeline, to a
point where it is answered by whoever is reading a log rather than by the person who owns the
deliverable.

WHAT THIS SCRIPT DOES AND DOES NOT DECIDE
-----------------------------------------
It MOVES BYTES. It reads the Kato records this run already has on disk, reads the model's
candidate rows for everything that is not on Kato, applies a mechanical duplicate sweep on
postal code, and writes the workbook.

It does NOT decide what a duplicate means, which email mention is a real option, or what
order the rows go in. Those are judgements and they arrive in master_candidates.json, which
the model writes after reading emails/emails.md and emails/_property_facts.json. Run this
with no candidates file and you get the Kato rows only, plus a loud warning: an inventory
that omits the email-only options is worse than no inventory, because it looks complete.

RE-RUNNING IS SAFE. If a workbook is already there, the user's Include? and Run notes are
carried forward by Row ID before it is rewritten, and the script says how many it preserved.
Pass --fresh to throw them away deliberately.
"""
import os, sys, re, json, argparse, datetime
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import load_config, read_json, write_json, display_name, norm_postcode
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

WORKBOOK = "Master List.xlsx"
MANIFEST = "master_list_manifest.json"
CANDIDATES = "master_candidates.json"
SHEET = "Master list"
SHEET_DUPES = "Duplicate check"
SHEET_FILES = "Unmatched files"

# CBRE chrome, taken from the reference template so the sheet a colleague opens is the one
# they have seen before. Green header band, amber for a row that overlaps another row, pale
# yellow for the two columns the user is being asked to fill in.
GREEN = "FF003F2D"
AMBER = "FFFFF2CC"
ASKED = "FFFFFFCC"
BAND = "FFF5F7F6"
# A row with no machine-readable document is the one thing on this sheet that costs the run
# something later: step 7a refuses an included row that has no PDF or PPTX to read, because
# with no document every specification field on its card comes from the tracker we generated
# ourselves, with no page-cited evidence behind it. "Link only" is not safer than "No" here -
# a link we did not download is bytes we do not hold. So anything that is not a flat Yes is
# filled red with bold black text, as real conditional formatting rather than a static fill,
# so it still reads correctly if the user sorts, filters or re-ranks the sheet.
NODOC = "FFFF8080"
FACE = "Arial"

HDR_ROW = 4
FIRST_ROW = 5
MIN_ROWS = 46          # keep the validated Include? dropdown on some spare rows below the data

# (header, key, width, number format, wrap)
COLUMNS = [
    ("Rank",                      "rank",             6,    "0",      False),
    ("Include?",                  "include",          11,   None,     False),
    ("Your Run notes for the AI", "run_notes",        26.3, None,     True),
    ("Property",                  "property",         34,   None,     True),
    ("Source type",               "source_type",      13,   None,     False),
    ("Source",                    "source",           46,   None,     True),
    ("Duplicate group",           "duplicate_group",  11,   None,     False),
    ("Duplicate status",          "duplicate_status", 26,   None,     True),
    ("Brochure?",                 "brochure",         12,   None,     False),
    ("Brochure detail",           "brochure_detail",  46,   None,     True),
    ("Address",                   "address",          34,   None,     True),
    ("Postcode",                  "postcode",         11,   None,     False),
    ("Latitude",                  "lat",              11,   "0.0000", False),
    ("Longitude",                 "lon",              11,   "0.0000", False),
    ("Size from (sq ft)",         "size_from",        13,   "#,##0",  False),
    ("Size to (sq ft)",           "size_to",          13,   "#,##0",  False),
    ("Size fit",                  "size_fit",         17,   None,     False),
    ("Tenure",                    "tenure",           13,   None,     False),
    ("Quoting rent / price",      "rent",             20,   None,     True),
    ("Availability status",       "availability",     24,   None,     True),
    ("Planning / scheme status",  "planning",         46,   None,     True),
    ("Available from",            "available_from",   13,   None,     False),
    ("Agent",                     "agent",            34,   None,     True),
    ("Agent email",               "agent_email",      30,   None,     False),
    ("Also offered by",           "also_offered_by",  34,   None,     True),
    ("Duplicate note",            "duplicate_note",   60,   None,     True),
    ("Notes from the source",     "notes",            80,   None,     True),
    # Hidden, and load-bearing. The user is free to sort, filter and re-rank the sheet, so
    # row position is not an identity. The Row ID travels with the row through any sort and
    # is what master_list_read.py keys every answer back onto. A row typed in by hand has no
    # Row ID, which is exactly how the read-back recognises one.
    ("Row ID",                    "row_id",           22,   None,     False),
]
KEYS = [c[1] for c in COLUMNS]
IDX = {k: i + 1 for i, (_h, k, *_r) in enumerate(COLUMNS)}
HIDDEN = {"row_id"}

DOC_EXTS = (".pdf", ".pptx")
FILE_EXTS = DOC_EXTS + (".xlsx", ".xlsm", ".csv", ".docx")
# Folders this run generates itself. A file in one of them is never a user upload.
GENERATED_DIRS = {"properties", "emails", "longlist_work", "longlist_inputs", "OUTPUT",
                  "_scratch", "plan_qa", "toolkit", "uploads_used"}
GENERATED_FILES = {WORKBOOK.lower(), "kato longlist (client).xlsx"}

INCLUDE_VALUES = ("Yes", "No")


# --------------------------------------------------------------------------- Kato side

def _tenure(d):
    if d.get("tenure"):
        return d["tenure"]
    parts = []
    if d.get("for_sale"):
        parts.append("For Sale")
    if d.get("to_let"):
        parts.append("To Let")
    return " / ".join(parts) or "tbd"


def _rent(d):
    rk = d.get("rent_kato") or {}
    s = (rk.get("string") or "").strip()
    if rk.get("from"):
        return "GBP%.2f psf" % float(rk["from"])
    if s and s.lower() not in ("-", "- non-quoting", "non-quoting", "poa"):
        return s
    pr = (d.get("price") or {}).get("string")
    if pr and d.get("for_sale"):
        return pr
    return "Rent on application"


def _brochure(pdir, d):
    """What document evidence this property actually has, stated as a fact about disk.

    Three states, and the middle one matters: a Kato listing regularly carries a brochure
    URL that the fetch could not retrieve, or a broker mails a link to a microsite rather
    than a PDF. Reporting that as 'No' hides a document the user can go and get; reporting
    it as 'Yes' promises the run bytes it does not hold.
    """
    media = os.path.join(pdir, "media")
    files = []
    if os.path.isdir(media):
        files = sorted(f for f in os.listdir(media)
                       if os.path.isfile(os.path.join(media, f))
                       and os.path.splitext(f)[1].lower() in DOC_EXTS)
    if files:
        kind = "PDF" if all(f.lower().endswith(".pdf") for f in files) else "file"
        return "Yes", "Yes - %d %s on disk: %s" % (len(files), kind, "; ".join(files))
    urls = [x for x in (d.get("documents") or []) if x.get("url")]
    if urls:
        names = "; ".join(x.get("name") or x.get("url") for x in urls[:3])
        return "Link only", "Link only - %d document link(s), not downloaded: %s" % (len(urls), names)
    return "No", "No document found for this property"


def _agents(d):
    names = [a.get("name") for a in (d.get("agents") or []) if a.get("name")]
    org = d.get("agent_organisation")
    label = "; ".join(names)
    if org:
        label = "%s (%s)" % (label, org) if label else org
    email = next((a.get("email") for a in (d.get("agents") or []) if a.get("email")), None)
    return label or "tbd", email or ""


def kato_rows(work, reqid):
    """One row per entry in properties/_index.json, Kato-sourced or materialised.

    A property materialised by an earlier pass of master_list_read.py is still in the index
    and still has a _derived.json, so it is picked up here exactly like a Kato one. It keeps
    the row id and source type it was created with, so a second build does not re-label a
    user's email option as a Kato listing and orphan their answer.
    """
    props_dir = os.path.join(work, "properties")
    idx = read_json(os.path.join(props_dir, "_index.json"), {}) or {}
    rows = []
    for p in idx.get("properties", []):
        folder = p.get("folder")
        pdir = os.path.join(props_dir, folder or "")
        d = read_json(os.path.join(pdir, "_derived.json"), {}) or {}
        if not d:
            continue
        a = d.get("address") or {}
        mp = (d.get("coordinates") or {}).get("map") or {}
        size = d.get("size") or {}
        broch, broch_detail = _brochure(pdir, d)
        agent, agent_email = _agents(d)
        origin = p.get("origin")
        if origin == "master_list":
            row_id = p.get("row_id") or "manual:%s" % folder
            src_type = p.get("source_type") or "Manual"
            src = p.get("source") or "Added from the master list"
        else:
            row_id = "kato:%s" % p.get("match_id")
            src_type = "Kato"
            src = "Kato requirement %s (Longlist), ID %s" % (reqid, p.get("match_id"))
        rows.append({
            "row_id": row_id,
            "folder": folder,
            "order": p.get("order"),
            "group_position": p.get("group_position"),
            "property": display_name({"address": a, "folder": folder}),
            "source_type": src_type,
            "source": src,
            "brochure": broch,
            "brochure_detail": broch_detail,
            "address": ", ".join(x for x in [a.get("line1"), a.get("line2"), a.get("town")] if x),
            "postcode": a.get("postcode") or "",
            "lat": mp.get("lat"),
            "lon": mp.get("lng"),
            "size_from": size.get("from"),
            "size_to": size.get("to") or size.get("from"),
            "tenure": _tenure(d),
            "rent": _rent(d),
            "availability": d.get("status") or "tbd",
            "planning": (d.get("summary") or "").strip(),
            "available_from": d.get("possession") or "tbd",
            "agent": agent,
            "agent_email": agent_email,
            "also_offered_by": "",
            "notes": (d.get("notes") or d.get("description") or "").strip(),
        })
    return rows, idx


# --------------------------------------------------------------- model candidate rows

def candidate_rows(cand, already):
    """`already` is the set of row ids the index already covers.

    A candidate the user accepted last pass has been MATERIALISED into properties/, so it
    comes back through the index with its own row id and its real folder. Emitting the
    candidate again as well would put the same option on the sheet twice, with only one of
    the two carrying the folder, and the duplicate sweep would then flag a property against
    itself. The index copy always wins: it is the one the run actually builds.
    """
    """Rows the model found that are not on Kato: email-only options and upload-only ones.

    Deliberately permissive about missing fields. An option first named in a one-line email
    ("we also have Packington Hill, 140k, Q2 2027") has no coordinates and no agent email,
    and refusing it would push exactly the options this step exists to surface back out of
    the inventory. Blanks show as blanks, and the row still carries a yes/no decision.
    """
    out, skipped = [], 0
    for i, r in enumerate(cand.get("rows") or [], 1):
        if not (r.get("property") or "").strip():
            print("  WARNING: candidate row %d has no 'property' name, skipped" % i, file=sys.stderr)
            continue
        rid = r.get("row_id") or "cand:%02d" % i
        if rid in already:
            skipped += 1
            continue
        out.append({
            "row_id": rid,
            "folder": None,
            "order": None,
            "group_position": None,
            "property": r.get("property").strip(),
            "source_type": r.get("source_type") or "Broker email",
            "source": r.get("source") or "",
            "brochure": r.get("brochure") or ("Yes" if r.get("files") else "No"),
            "brochure_detail": r.get("brochure_detail") or (
                "Yes - %d file(s) supplied: %s" % (len(r["files"]), "; ".join(r["files"]))
                if r.get("files") else "No document identified for this property"),
            "address": r.get("address") or "",
            "postcode": r.get("postcode") or "",
            "lat": r.get("lat"),
            "lon": r.get("lon"),
            "size_from": r.get("size_from"),
            "size_to": r.get("size_to") if r.get("size_to") is not None else r.get("size_from"),
            "tenure": r.get("tenure") or "tbd",
            "rent": r.get("rent") or "tbd",
            "availability": r.get("availability") or "tbd",
            "planning": r.get("planning") or "tbd",
            "available_from": r.get("available_from") or "tbd",
            "agent": r.get("agent") or "tbd",
            "agent_email": r.get("agent_email") or "",
            "also_offered_by": r.get("also_offered_by") or "",
            "notes": r.get("notes") or "",
            "files": r.get("files") or [],
        })
    if skipped:
        print("  %d candidate row(s) already materialised into properties/, taken from the index"
              % skipped, flush=True)
    return out


# ------------------------------------------------------------------------- judgements

def apply_size_fit(rows, brief):
    """'In brief' / 'Below brief' / 'Above brief', from the requirement's own size band.

    Left blank rather than guessed when the model supplied no band: a column that says
    'In brief' on every row because there was nothing to compare against is a column that
    makes a user stop checking.
    """
    lo = (brief or {}).get("size_min")
    hi = (brief or {}).get("size_max")
    for r in rows:
        f, t = r.get("size_from"), r.get("size_to")
        if f is None and t is None:
            r["size_fit"] = "Not stated"
            continue
        if lo is None and hi is None:
            r["size_fit"] = ""
            continue
        f = f if f is not None else t
        t = t if t is not None else f
        if hi is not None and f > hi:
            r["size_fit"] = "Above brief"
        elif lo is not None and t < lo:
            r["size_fit"] = "Below brief"
        else:
            r["size_fit"] = "In brief"


def apply_duplicates(rows, cand):
    """Group rows that may be the same building, explicit groups first, then a postcode sweep.

    The explicit groups are the model's adjudication and they win: it has read the emails
    and can tell a phase from a re-listing. The sweep then catches what is left, on postal
    code alone, and is deliberately blunt. It over-groups (a large park shares one code
    across genuinely different units) and it says so in the status text, because the cost of
    a group the user glances at and dismisses is a second of their time, and the cost of a
    missed duplicate is the same building on the client's dashboard twice.

    Rows with no postal code are never swept. An empty string is not evidence of anything.

    Each group is recorded with its ORIGIN and whether it is MERGEABLE, and the two are not
    the same question. Mergeable decides what happens to a row the user rejects inside the
    group: merged into the survivor, so its brochure travels with it, or simply dropped. A
    group the model adjudicated is mergeable because somebody read the sources and said these
    are one building. An automatic group is mergeable only when the floor areas are identical,
    because a shared postal code on its own routinely covers genuinely different units on one
    park, and handing one unit's brochure to another unit's card is worse than losing it.
    """
    by_id = {r["row_id"]: r for r in rows}
    meta = {}
    n = 0
    for gid, g in sorted((cand.get("duplicate_groups") or {}).items()):
        members = [by_id[m] for m in (g.get("members") or []) if m in by_id]
        missing = [m for m in (g.get("members") or []) if m not in by_id]
        if missing:
            print("  WARNING: duplicate group %s names unknown row id(s): %s"
                  % (gid, ", ".join(missing)), file=sys.stderr)
        if len(members) < 2:
            continue
        n = max(n, _gnum(gid))
        status = g.get("status") or "Flagged by the model as the same option"
        meta[gid] = {"origin": "model", "mergeable": True, "status": status,
                     "members": [r["row_id"] for r in members]}
        for r in members:
            r["duplicate_group"] = gid
            r["duplicate_status"] = status
            r["duplicate_note"] = g.get("note") or ""

    buckets = {}
    for r in rows:
        if r.get("duplicate_group"):
            continue
        pc = norm_postcode(r.get("postcode"))
        if not pc:
            continue
        buckets.setdefault(pc, []).append(r)
    for pc, group in sorted(buckets.items()):
        if len(group) < 2:
            continue
        n += 1
        gid = "D%d" % n
        sizes = {r.get("size_to") or r.get("size_from") for r in group}
        same_size = len(sizes) == 1 and None not in sizes
        status = ("SAME BUILDING? same postcode and floor area (auto)" if same_size
                  else "Same postcode, sizes differ (auto)")
        note = ("Automatic check on postal code %s only, not adjudicated: %s. %s"
                % (pc, "; ".join(r["property"] for r in group),
                   "Identical floor area, so most likely one building listed more than once."
                   if same_size else
                   "Could be separate units on one park, or one building quoted at different sizes."))
        meta[gid] = {"origin": "auto", "mergeable": bool(same_size), "status": status,
                     "members": [r["row_id"] for r in group]}
        for r in group:
            r["duplicate_group"] = gid
            r["duplicate_status"] = status
            r["duplicate_note"] = note
    for r in rows:
        r.setdefault("duplicate_group", "")
        r.setdefault("duplicate_status", "")
        r.setdefault("duplicate_note", "")
    return meta


def _gnum(gid):
    m = re.search(r"(\d+)$", str(gid) or "")
    return int(m.group(1)) if m else 0


def apply_rank(rows, cand):
    """Order the sheet. The model's explicit order first, then everything it did not rank.

    Unranked rows keep the Kato longlist's own order (group_position, then fetch order) and
    the candidate rows follow, because that is a stable order the user can predict, and a
    sheet whose rows move on every rebuild is a sheet whose answers cannot be trusted.
    """
    order = [str(x) for x in (cand.get("rank") or [])]
    pos = {rid: i for i, rid in enumerate(order)}
    big = len(pos) + 1

    def key(r):
        if r["row_id"] in pos:
            return (0, pos[r["row_id"]], 0, "")
        gp = r.get("group_position")
        od = r.get("order")
        if od is not None:
            return (1, big, gp if gp is not None else od, "")
        return (2, big, 0, r["property"].lower())

    rows.sort(key=key)
    for i, r in enumerate(rows, 1):
        r["rank"] = i


def carry_forward(path, rows):
    """Preserve the user's existing answers across a rebuild, keyed on Row ID.

    A rebuild happens for a real reason (a second email export landed, the user dropped in
    three more brochures) and it must not cost them the forty decisions they already made.
    Answers are matched on the hidden Row ID first, because that is an identity and survives
    any sort. A row the user TYPED IN has no Row ID, so those fall back to property name plus
    postal code: master_list_read.py materialises a hand-typed row and gives it a Row ID in
    the index, but it cannot put that id back in the workbook without re-saving a file the
    user may have open. The fallback is only ever consulted for a prior row that had no Row
    ID, so it can never override an identity match.
    """
    if not os.path.exists(path):
        return 0, {}
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as e:
        print("  WARNING: could not read the existing workbook to carry answers forward (%s)"
              % e, file=sys.stderr)
        return 0, {}
    if SHEET not in wb.sheetnames:
        return 0, {}
    ws = wb[SHEET]
    hdr = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(HDR_ROW, c).value
        if v:
            hdr[str(v).strip()] = c
    need = ("Row ID", "Include?", "Your Run notes for the AI")
    if not all(h in hdr for h in need):
        return 0, {}
    def nk(name, pc):
        return (re.sub(r"\s+", " ", str(name or "").strip().lower()),
                re.sub(r"\s+", "", str(pc or "").strip().lower()))

    prior, by_name, manual = {}, {}, {}
    for r in range(FIRST_ROW, ws.max_row + 1):
        rid = ws.cell(r, hdr["Row ID"]).value
        inc = ws.cell(r, hdr["Include?"]).value
        note = ws.cell(r, hdr["Your Run notes for the AI"]).value
        name = ws.cell(r, hdr["Property"]).value if "Property" in hdr else None
        pc = ws.cell(r, hdr["Postcode"]).value if "Postcode" in hdr else None
        if not rid and not (inc or note):
            continue
        if not rid:
            manual[r] = {h: ws.cell(r, c).value for h, c in hdr.items()}
            if name and (inc or note):
                by_name[nk(name, pc)] = (inc, note)
            continue
        if inc or note:
            prior[str(rid).strip()] = (inc, note)
    kept = 0
    matched_names = set()
    for row in rows:
        got = prior.get(row["row_id"])
        if got is None:
            key = nk(row.get("property"), row.get("postcode"))
            got = by_name.get(key)
            if got is not None:
                matched_names.add(key)
        if got:
            row["include"] = got[0] or ""
            row["run_notes"] = got[1] or ""
            kept += 1
    # A hand-typed row is only "lost" if nothing re-matched it. One that was materialised
    # last pass comes back through the index and matches on name, so reporting it as lost
    # would send the user hunting for a row that is sitting in front of them.
    for r, rec in list(manual.items()):
        if nk(rec.get("Property"), rec.get("Postcode")) in matched_names:
            del manual[r]
    return kept, manual


# ------------------------------------------------------------------------- the workbook

def _hdr_cell(ws, r, c, text):
    cell = ws.cell(r, c, text)
    cell.font = Font(name=FACE, sz=9, bold=True, color="FFFFFFFF")
    cell.fill = PatternFill("solid", fgColor=GREEN)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return cell


def _flag_missing_documents(ws, last_data_row):
    """Red fill + bold black on every row whose Brochure? is not a flat Yes.

    Driven off the Brochure? column, whose three values are stated by _brochure() as facts
    about disk: Yes (a PDF or PPTX is here), Link only (a URL nobody downloaded) and No.
    Only Yes is left alone. Blank is excluded from the rule so the spare rows below the data,
    which carry the Include? dropdown and no property, are not painted red.
    """
    if last_data_row < FIRST_ROW:
        return
    broch = get_column_letter(IDX["brochure"])
    # The Brochure? column and nothing else. This is a pre-delivery flag for the person
    # scanning the sheet to spot the rows with no brochure, so it belongs on the one cell
    # that states the answer; spreading it across neighbouring columns only makes the sheet
    # noisier without telling the reader anything more.
    rng = "%s%d:%s%d" % (broch, FIRST_ROW, broch, last_data_row)
    rule = FormulaRule(
        formula=['AND($%s%d<>"",$%s%d<>"Yes")' % (broch, FIRST_ROW, broch, FIRST_ROW)],
        fill=PatternFill("solid", start_color=NODOC, end_color=NODOC),
        font=Font(name=FACE, sz=9, bold=True, color="FF000000"),
        stopIfTrue=False)
    ws.conditional_formatting.add(rng, rule)


def write_master(path, rows, title, note, footnotes):
    thin = Side(style="thin", color="FFD9D9D9")
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.sheet_view.showGridLines = False

    t = ws.cell(1, 1, title)
    t.font = Font(name=FACE, sz=14, bold=True, color=GREEN)
    ws.row_dimensions[1].height = 18
    n = ws.cell(2, 1, note)
    n.font = Font(name=FACE, sz=9)

    for i, (h, key, w, _f, _wrap) in enumerate(COLUMNS, 1):
        _hdr_cell(ws, HDR_ROW, i, h)
        ws.column_dimensions[get_column_letter(i)].width = w
        if key in HIDDEN:
            ws.column_dimensions[get_column_letter(i)].hidden = True
    ws.row_dimensions[HDR_ROW].height = 33.95

    band = False
    for j, row in enumerate(rows):
        r = FIRST_ROW + j
        ws.row_dimensions[r].height = 45.95
        dup = bool(row.get("duplicate_group"))
        if not dup:
            band = not band
        base = AMBER if dup else (BAND if band else None)
        for i, (_h, key, _w, fmt, wrap) in enumerate(COLUMNS, 1):
            cell = ws.cell(r, i, row.get(key) if row.get(key) not in ("",) else None)
            cell.font = Font(name=FACE, sz=9)
            cell.border = Border(bottom=thin)
            if fmt:
                cell.number_format = fmt
            cell.alignment = Alignment(vertical="top", wrap_text=bool(wrap))
            fill = ASKED if key in ("include", "run_notes") else base
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)

    last = FIRST_ROW + max(len(rows), MIN_ROWS) - 1
    for r in range(FIRST_ROW + len(rows), last + 1):
        ws.row_dimensions[r].height = 45.95
        for key in ("include", "run_notes"):
            c = ws.cell(r, IDX[key])
            c.font = Font(name=FACE, sz=9)
            c.fill = PatternFill("solid", fgColor=ASKED)
            c.border = Border(bottom=thin)
            c.alignment = Alignment(vertical="top", wrap_text=(key == "run_notes"))

    _flag_missing_documents(ws, FIRST_ROW + len(rows) - 1)

    dv = DataValidation(type="list", formula1='"%s"' % ",".join(INCLUDE_VALUES),
                        allow_blank=True, showDropDown=False)
    dv.prompt = "Choose Yes or No"
    dv.promptTitle = "Include this property in the run?"
    ws.add_data_validation(dv)
    col = get_column_letter(IDX["include"])
    dv.add("%s%d:%s%d" % (col, FIRST_ROW, col, last))

    ws.auto_filter.ref = "A%d:%s%d" % (HDR_ROW, get_column_letter(len(COLUMNS)), last)
    ws.freeze_panes = "D%d" % FIRST_ROW

    r = last + 2
    for line in footnotes:
        c = ws.cell(r, 1, line)
        c.font = Font(name=FACE, sz=8, color="FF666666")
        r += 1
    return wb


def write_dupes(wb, rows):
    """The tab row 2 points at: every flagged group, its members side by side, nothing else.

    Separate from the master list because the decision it supports is a comparison between
    two rows, and on a twenty-seven column sheet sorted by rank those two rows are rarely
    adjacent. Read-only: the answer still goes in the Include? column on the master list, so
    there is exactly one place a decision lives.
    """
    ws = wb.create_sheet(SHEET_DUPES)
    ws.sheet_view.showGridLines = False
    groups = {}
    for r in rows:
        if r.get("duplicate_group"):
            groups.setdefault(r["duplicate_group"], []).append(r)

    t = ws.cell(1, 1, "Duplicate check")
    t.font = Font(name=FACE, sz=14, bold=True, color=GREEN)
    sub = ("Each block below is one set of rows that may be the same building. Decide which one "
           "to keep, then set Include? on the Master list tab. Nothing on this tab is read by the run.")
    ws.cell(2, 1, sub).font = Font(name=FACE, sz=9)
    if not groups:
        ws.cell(4, 1, "No overlapping rows found.").font = Font(name=FACE, sz=9)
        for letter, w in zip("ABCDEFGH", (12, 40, 14, 14, 14, 20, 34, 70)):
            ws.column_dimensions[letter].width = w
        return ws

    cols = [("Rank", "rank", 8), ("Property", "property", 40), ("Source type", "source_type", 14),
            ("Postcode", "postcode", 12), ("Size to (sq ft)", "size_to", 14),
            ("Quoting rent / price", "rent", 20), ("Agent", "agent", 34),
            ("Brochure detail", "brochure_detail", 46), ("Duplicate note", "duplicate_note", 70)]
    for i, (h, _k, w) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    r = 4
    for gid in sorted(groups, key=_gnum):
        members = sorted(groups[gid], key=lambda x: x.get("rank") or 0)
        head = ws.cell(r, 1, "%s - %s" % (gid, members[0].get("duplicate_status") or "possible duplicate"))
        head.font = Font(name=FACE, sz=10, bold=True, color=GREEN)
        r += 1
        for i, (h, _k, _w) in enumerate(cols, 1):
            _hdr_cell(ws, r, i, h)
        r += 1
        for m in members:
            for i, (_h, k, _w) in enumerate(cols, 1):
                c = ws.cell(r, i, m.get(k) if m.get(k) not in ("",) else None)
                c.font = Font(name=FACE, sz=9)
                c.fill = PatternFill("solid", fgColor=AMBER)
                c.alignment = Alignment(vertical="top", wrap_text=(k in ("property", "brochure_detail",
                                                                        "duplicate_note", "agent")))
            ws.row_dimensions[r].height = 40
            r += 1
        r += 1
    return ws


def scan_uploads(work, rows):
    """Files the user put in the working directory that no row claims.

    A brochure nobody tied to a property does nothing at all in this run. It is not an error
    and it is not necessarily a mistake (a market report, a floor plan for an option that was
    dropped), but it has to be VISIBLE, or the user is left believing a file they supplied was
    read. Only the working directory root and an uploads/ folder are scanned: everything the
    run generates lives in named folders and is skipped by name.
    """
    claimed = set()
    for r in rows:
        for f in (r.get("files") or []):
            claimed.add(os.path.normcase(os.path.basename(str(f))))
    # A file already copied into a property's media/ is claimed whatever the rows now say.
    # After the first accept-and-materialise pass the row no longer carries the filename,
    # because it comes back through the index, and reporting the original in uploads/ as
    # unmatched would tell the user their brochure was ignored when the run is using it.
    props_dir = os.path.join(work, "properties")
    if os.path.isdir(props_dir):
        for folder in os.listdir(props_dir):
            media = os.path.join(props_dir, folder, "media")
            if os.path.isdir(media):
                for f in os.listdir(media):
                    claimed.add(os.path.normcase(f))
    found = []
    roots = [work, os.path.join(work, "uploads")]
    for root in roots:
        if not os.path.isdir(root):
            continue
        for name in sorted(os.listdir(root)):
            p = os.path.join(root, name)
            if os.path.isdir(p):
                continue
            if os.path.splitext(name)[1].lower() not in FILE_EXTS:
                continue
            if name.lower() in GENERATED_FILES or name.startswith("~$"):
                continue
            if os.path.normcase(name) in claimed:
                continue
            st = os.stat(p)
            found.append((os.path.relpath(p, work), round(st.st_size / 1024.0),
                          datetime.datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d %H:%M")))
    return found


def write_files_sheet(wb, found):
    ws = wb.create_sheet(SHEET_FILES)
    ws.sheet_view.showGridLines = False
    ws.cell(1, 1, "Unmatched files").font = Font(name=FACE, sz=14, bold=True, color=GREEN)
    ws.cell(2, 1, "Files in the working directory that no row on the Master list claims. They will "
                  "not be read by the run. Tell the run which property each belongs to, or ignore "
                  "them.").font = Font(name=FACE, sz=9)
    for letter, w in zip("ABC", (70, 14, 20)):
        ws.column_dimensions[letter].width = w
    for i, h in enumerate(("File", "Size (KB)", "Modified"), 1):
        _hdr_cell(ws, 4, i, h)
    if not found:
        ws.cell(5, 1, "None. Every file in the working directory is accounted for.").font = Font(name=FACE, sz=9)
        return ws
    for j, (rel, kb, mtime) in enumerate(found):
        r = 5 + j
        for i, v in enumerate((rel, kb, mtime), 1):
            c = ws.cell(r, i, v)
            c.font = Font(name=FACE, sz=9)
            if i == 2:
                c.number_format = "#,##0"
    return ws


# ------------------------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", required=True)
    ap.add_argument("--candidates", default=None,
                    help="model-written candidate rows, duplicate groups, rank and brief "
                         "(default: <work>/%s)" % CANDIDATES)
    ap.add_argument("--out", default=None, help="output workbook (default: <work>/%s)" % WORKBOOK)
    ap.add_argument("--fresh", action="store_true",
                    help="discard any Include?/Run notes already in the workbook instead of "
                         "carrying them forward by Row ID")
    args = ap.parse_args()

    cfg = load_config(args.config)
    work = cfg["work_dir"]
    out = args.out or os.path.join(work, WORKBOOK)
    cand_path = args.candidates or os.path.join(work, CANDIDATES)
    cand = read_json(cand_path, {}) or {}
    if not os.path.exists(cand_path):
        print("WARNING: no %s, so this inventory covers the Kato longlist ONLY.\n"
              "  Every option that exists only in a broker email or an uploaded file is MISSING "
              "from it,\n  and a user reading it will reasonably assume it is the complete "
              "picture. Write the\n  candidates file from emails/emails.md + "
              "emails/_property_facts.json and re-run."
              % os.path.relpath(cand_path, work), file=sys.stderr)

    props_dir = os.path.join(work, "properties")
    if not os.path.exists(os.path.join(props_dir, "_index.json")):
        raise SystemExit("No properties/_index.json in %s - run stage 1 (or 1-alt) first." % work)

    idx_reqid = (read_json(os.path.join(props_dir, "_index.json"), {}) or {}).get("requirement_id")
    rows, _idx = kato_rows(work, idx_reqid)
    n_kato = sum(1 for r in rows if r["source_type"] == "Kato")
    n_index = len(rows)
    rows += candidate_rows(cand, {r["row_id"] for r in rows})

    apply_size_fit(rows, cand.get("brief"))
    dup_meta = apply_duplicates(rows, cand)
    apply_rank(rows, cand)
    for r in rows:
        r.setdefault("include", "")
        r.setdefault("run_notes", "")

    kept, manual = (0, {}) if args.fresh else carry_forward(out, rows)

    client = cfg.get("client") or "this requirement"
    title = cand.get("title") or ("Master list - %s - every option found across Kato, the broker "
                                  "emails and the uploaded files" % client)
    note = ("Set the Include? column to Yes or No on every row, and put any instruction for the "
            "run in 'Your Run notes for the AI'. Amber rows overlap with another row - see the "
            "Duplicate check tab.")
    basis = cand.get("rank_basis")
    foot = ["Rank is a sort order%s. Re-sort or re-rank freely: the run reads your answers by a "
            "hidden row id, not by position." % (": %s" % basis if basis else ""),
            "Sources: %d Kato longlist record(s); %d option(s) named only in the broker emails or "
            "the uploaded files. Built %s."
            % (n_kato, len(rows) - n_kato, datetime.datetime.now().strftime("%Y-%m-%d %H:%M")),
            "Red, bold Brochure? cells are rows with no machine-readable document: 'No' means "
            "nothing was found, 'Link only' means a URL nobody downloaded, and neither gives the "
            "run a page to cite. Include one and its card's specification comes from our own "
            "tracker with no evidence behind it, so either chase the brochure or expect the row "
            "in the Gaps Report. Only a flat 'Yes' is left unformatted.",
            "Nothing here is sent to a client. It decides what the run builds."]

    wb = write_master(out, rows, title, note, foot)
    write_dupes(wb, rows)
    unmatched = scan_uploads(work, rows)
    write_files_sheet(wb, unmatched)

    try:
        wb.save(out)
    except PermissionError:
        raise SystemExit("Cannot write %s - it is open in Excel. Close it and re-run." % out)

    manifest = {"generated": datetime.datetime.now().isoformat(timespec="seconds"),
                "workbook": os.path.basename(out),
                "requirement_id": idx_reqid,
                "kato_rows": n_kato,
                "index_rows": n_index,
                "candidate_rows": len(rows) - n_index,
                "unmatched_files": [u[0] for u in unmatched],
                # Per group: where it came from and whether a rejected member may be merged
                # into the kept one. master_list_read.py reads "mergeable" and nothing else.
                "duplicate_groups": dup_meta,
                "rows": [{k: r.get(k) for k in
                          ("row_id", "rank", "property", "source_type", "source", "folder",
                           "postcode", "duplicate_group", "duplicate_status", "files")}
                         for r in rows]}
    write_json(os.path.join(work, MANIFEST), manifest)

    dups = len({r["duplicate_group"] for r in rows if r.get("duplicate_group")})
    print("MASTER LIST -> %s" % out, flush=True)
    print("  rows=%d (kato=%d, already materialised=%d, new from emails/uploads=%d)  "
          "duplicate groups=%d  unmatched files=%d"
          % (len(rows), n_kato, n_index - n_kato, len(rows) - n_index, dups, len(unmatched)),
          flush=True)
    if kept:
        print("  carried forward %d existing answer(s) by Row ID" % kept, flush=True)
    if manual:
        print("  NOTE: %d hand-typed row(s) in the previous workbook had no Row ID and were NOT "
              "carried forward. Re-add them: %s"
              % (len(manual), "; ".join(str((v.get("Property") or "?")) for v in manual.values())),
              file=sys.stderr)
    if unmatched:
        print("  %d file(s) in the working directory belong to no row - see the '%s' tab"
              % (len(unmatched), SHEET_FILES), flush=True)
    print("  NEXT: the user fills in Include? and Run notes, then run master_list_read.py.", flush=True)


if __name__ == "__main__":
    main()
