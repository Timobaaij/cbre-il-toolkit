#!/usr/bin/env python3
"""
Stage 2.5c - read the user's answers off the MASTER LIST and make them binding.

Turns the edited workbook into master_list.json, the single source of truth every later
stage reads, and materialises a property folder for each included option that exists only
in a broker email or an uploaded file, so the rest of the pipeline can treat it exactly
like a Kato record.

WHAT "BINDING" MEANS HERE
-------------------------
Before this step the run's shape was decided by three separate mechanisms that never saw
each other: whatever Kato returned, an automatic merge on postal code and floor area, and
nothing at all for email-only options. This file replaces all three with one answer per
property, given by the person responsible for the deliverable.

  * Include? = No          the property is not built. If it shares a duplicate group with
                           an included row it becomes that row's merged sibling instead of
                           simply vanishing, so its brochure still reaches the pipeline
                           (see merge_map below).
  * Include? = Yes         the property is built. If it has no folder yet, one is created.
  * Your Run notes         free text, carried to the model at enrichment and to the Gaps
                           Report. Instructions like "split this unit into three cards" or
                           "run but without rent" are acted on by the model, not by this
                           script: it moves the text, it does not interpret it.

YES OR NO, AND NOTHING ELSE
---------------------------
There is no third value and anything that is not Yes or No stops the run, "maybe" included.
A deferred answer still has to be resolved before the run can start, so accepting one here
would not save the round trip, it would only move the decision to a point where it is taken
by whoever is reading a log instead of by the person who owns the deliverable. The refusal
names every offending row, so resolving them is one pass through the sheet.

WHY A MISSING ROW IS NOT AN ANSWER
----------------------------------
A row in the manifest that is not in the workbook was deleted by the user. That reads like
"no", but it is indistinguishable from a botched sort or a filtered copy saved over the
original, so it is reported and treated as excluded rather than acted on quietly. A row in
the workbook with no Row ID is the opposite case and is fully supported: the user typed in
an option the run never found, and if they marked it Yes it is materialised like any other.
"""
import os, sys, re, json, shutil, argparse, datetime
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import load_config, read_json, write_json, property_folder, sanitize
from openpyxl import load_workbook

import master_list_build as B

MASTER_JSON = "master_list.json"
YES, NO = "Yes", "No"
_TRUE = {"yes": YES, "y": YES, "include": YES, "true": YES, "1": YES}
_FALSE = {"no": NO, "n": NO, "exclude": NO, "drop": NO, "false": NO, "0": NO}
# Deliberately NOT a value. Mapped only so the refusal can say "maybe is not an answer here"
# rather than the unhelpful "unreadable value 'Maybe'".
_DEFERRED = {"maybe", "m", "tbc", "tbd", "?", "unsure", "not sure"}


def find_workbook(work, explicit):
    """The workbook, even when the user saved it under another name.

    Same contract as stage 2's hunt for the email export: look for the expected name, then
    for any workbook in the working directory that IS a master list (sheet name and header
    cell), and SAY which file was used. A user who saved 'Master List (Toby edits).xlsx' and
    got a run built off the untouched original would have no way of telling from the output.
    """
    if explicit:
        if not os.path.exists(explicit):
            raise SystemExit("No such workbook: %s" % explicit)
        return explicit
    direct = os.path.join(work, B.WORKBOOK)
    if os.path.exists(direct):
        return direct
    cands = []
    for name in sorted(os.listdir(work)):
        if not name.lower().endswith((".xlsx", ".xlsm")) or name.startswith("~$"):
            continue
        p = os.path.join(work, name)
        try:
            wb = load_workbook(p, read_only=True, data_only=True)
            if B.SHEET in wb.sheetnames and str(wb[B.SHEET].cell(B.HDR_ROW, 1).value or "").strip() == "Rank":
                cands.append(p)
            wb.close()
        except Exception:
            continue
    if not cands:
        raise SystemExit("No master list workbook found in %s. Run master_list_build.py first." % work)
    if len(cands) > 1:
        print("WARNING: %d master list workbooks in the working directory; using the most recently "
              "modified. The others: %s"
              % (len(cands), "; ".join(os.path.basename(c) for c in cands)), file=sys.stderr)
        cands.sort(key=os.path.getmtime, reverse=True)
    return cands[0]


def read_rows(path):
    wb = load_workbook(path, data_only=True)
    if B.SHEET not in wb.sheetnames:
        raise SystemExit("%s has no '%s' sheet." % (path, B.SHEET))
    ws = wb[B.SHEET]
    hdr = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(B.HDR_ROW, c).value
        if v:
            hdr[str(v).strip()] = c
    header_to_key = {h: k for (h, k, *_r) in B.COLUMNS}
    missing = [h for h in ("Include?", "Property") if h not in hdr]
    if missing:
        raise SystemExit("%s is missing required column(s): %s" % (path, ", ".join(missing)))
    rows = []
    for r in range(B.FIRST_ROW, ws.max_row + 1):
        rec = {}
        for h, c in hdr.items():
            key = header_to_key.get(h)
            if key:
                rec[key] = ws.cell(r, c).value
        rec["_excel_row"] = r
        if not (rec.get("property") or rec.get("row_id") or rec.get("include")):
            continue
        for k, v in list(rec.items()):
            if isinstance(v, str):
                rec[k] = v.strip()
        rows.append(rec)
    wb.close()
    return rows


def norm_include(v):
    if v is None:
        return ""
    s = str(v).strip()
    low = s.lower()
    return _TRUE.get(low) or _FALSE.get(low) or s


def build_merge_map(rows, dup_meta):
    """Within one duplicate group: the excluded rows become the included row's siblings.

    The whole point of adjudicating duplicates is that the user picks WHICH listing of a
    building to ship, and the listings are rarely equally complete. One broker writes the
    fuller description, another attaches the only brochure. Dropping the rejected listing
    outright throws the brochure away and the surviving card ships unevidenced, which is the
    defect this repository already has a guard for (toolkit_tracker's unevidenced-row
    refusal). So a rejected row inside a group is merged into the kept one rather than
    deleted, and its documents travel with it.

    Only when exactly ONE row in the group is included. Two included rows means the user
    looked and decided they are different buildings, and nothing here second-guesses that.
    Zero included rows means the whole group was rejected and there is nothing to merge into.

    And only when the GROUP IS MERGEABLE, which the build step recorded per group: a group
    the model adjudicated, or an automatic one whose members share an exact floor area. An
    automatic group formed on postal code alone with differing areas is usually separate
    units on one park, and merging there would hand one unit's brochure to another unit's
    card, which is a wrong page citation rather than a missing one. Those are excluded
    plainly and reported, so a lost brochure is visible instead of silently misfiled.
    """
    groups = {}
    for r in rows:
        g = (r.get("duplicate_group") or "").strip()
        if g:
            groups.setdefault(g, []).append(r)
    merge, kept_separate, not_merged = {}, [], []
    for g, members in sorted(groups.items()):
        inc = [m for m in members if m["include"] == YES]
        exc = [m for m in members if m["include"] == NO]
        meta = dup_meta.get(g) or {}
        # Unknown group (the user typed a group id of their own): fall back to the same
        # evidence the sweep uses, an identical floor area on the sheet itself.
        if meta:
            mergeable = bool(meta.get("mergeable"))
        else:
            sizes = {m.get("size_to") or m.get("size_from") for m in members}
            mergeable = len(sizes) == 1 and None not in sizes
        if len(inc) == 1 and exc:
            tgt = inc[0]
            for m in exc:
                if not m.get("folder"):
                    continue
                if mergeable:
                    merge[m["folder"]] = {"into": tgt.get("folder"), "into_row_id": tgt.get("row_id"),
                                          "group": g, "property": m.get("property")}
                else:
                    not_merged.append((g, m.get("property"), tgt.get("property")))
        elif len(inc) > 1:
            kept_separate.append((g, [m.get("property") for m in inc]))
    return merge, kept_separate, not_merged


# --------------------------------------------------------------- materialisation

def _num(v):
    if v in (None, "", "tbd"):
        return None
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def _int(v):
    f = _num(v)
    return int(round(f)) if f is not None else None


def synth_derived(row):
    """A _derived.json for an option that was never on Kato.

    Same keys as common.derive() produces, because every later stage reads that shape and a
    partial one would fail somewhere far from here with an unhelpful message. Everything the
    master list does not carry is None or empty, never invented: the run's own honesty gates
    are downstream of this file and they can only report a gap they can see.
    """
    pc = (row.get("postcode") or "").strip()
    addr_line = (row.get("address") or "").strip()
    segs = [s.strip() for s in addr_line.split(",") if s.strip()]
    town = segs[-1] if segs else None
    lat, lon = _num(row.get("lat")), _num(row.get("lon"))
    sf, st = _int(row.get("size_from")), _int(row.get("size_to"))
    tenure = (row.get("tenure") or "").strip()
    rent = (row.get("rent") or "").strip()
    agent = (row.get("agent") or "").strip()
    email = (row.get("agent_email") or "").strip()
    org = None
    m = re.search(r"\(([^)]+)\)\s*$", agent)
    if m:
        org = m.group(1)
    elif "," in agent:
        org = agent.rsplit(",", 1)[1].strip()
    notes = (row.get("notes") or "").strip()
    planning = (row.get("planning") or "").strip()
    return {
        "match_id": None,
        "status": (row.get("availability") or "tbd"),
        "to_let": "let" in tenure.lower() or not tenure,
        "for_sale": "sale" in tenure.lower(),
        "tenure": tenure or None,
        "possession": (row.get("available_from") or None),
        "address": {"name": row.get("property"), "line1": segs[0] if segs else None,
                    "line2": segs[1] if len(segs) > 2 else None, "town": town, "county": None,
                    "postcode": pc or None, "uprn": None,
                    "full": ", ".join(x for x in [row.get("property"), addr_line, pc] if x)},
        "area": None,
        "coordinates": {"map": ({"lat": lat, "lng": lon} if lat is not None and lon is not None else None),
                        "street_view": None},
        "size": {"from": sf, "to": st, "string": ("%s sq ft" % format(st or sf, ",")) if (st or sf) else None},
        # The workbook fills an unstated rent with "tbd". Carrying that through as a rent
        # would put the placeholder itself on the card, so only a real value is written.
        "rent_kato": {"string": (rent if (rent or "").strip().lower() not in
                                 ("", "tbd", "n/a", "na", "not stated", "none") else None),
                      "from": None, "to": None},
        "price": {"string": None, "value": None},
        "service_charge": None, "rates_payable": None, "estate_charge": None,
        "total_sqft": st or sf, "total_pa": None, "epc": None, "lease": None,
        "building_types": None, "fitted_space": None,
        "key_points": [x for x in [planning if planning and planning != "tbd" else None] if x],
        "amenities": [],
        "summary": planning if planning and planning != "tbd" else None,
        "description": notes or None,
        "location_text": None,
        "notes": notes or None,
        "website": None, "videos": [], "documents": [], "images": [],
        "available_spaces": [], "tube": None, "train": None,
        "agents": ([{"name": agent, "position": None, "tel": None, "mobile": None, "email": email or None}]
                   if agent and agent != "tbd" else []),
        "agent_organisation": org,
        "landlord_confidential": False, "landlord_companies": [],
        "messages": [],
        "published_at": None, "updated_at": None, "group_position": None,
        "_origin": {"source": "master list", "source_type": row.get("source_type"),
                    "source_detail": row.get("source"), "row_id": row.get("row_id"),
                    "created": datetime.datetime.now().isoformat(timespec="seconds")},
    }


def resolve_file(work, ref):
    """Find a file the master list names, wherever the user actually put it."""
    ref = str(ref).strip().strip('"')
    for cand in (os.path.join(work, ref), ref):
        if os.path.isfile(cand):
            return cand
    base = os.path.basename(ref)
    for root in (work, os.path.join(work, "uploads")):
        p = os.path.join(root, base)
        if os.path.isfile(p):
            return p
    for root, dirs, files in os.walk(work):
        dirs[:] = [d for d in dirs if d not in B.GENERATED_DIRS and not d.startswith(".")]
        if base in files:
            return os.path.join(root, base)
    return None


def materialise(work, rows, manifest_files):
    """Create properties/<folder>/ for every included row that has none yet.

    Without this the master list is decoration for exactly the options it exists to rescue:
    the run's spine is properties/_index.json, so an email-only option marked Yes would be
    read back, counted, reported, and then built by nothing. Idempotent on row_id, so a
    second pass updates in place rather than creating a second folder.
    """
    props_dir = os.path.join(work, "properties")
    idx_path = os.path.join(props_dir, "_index.json")
    idx = read_json(idx_path, {}) or {}
    entries = idx.get("properties", [])
    by_row = {e.get("row_id"): e for e in entries if e.get("row_id")}
    next_order = max([int(e.get("order") or 0) for e in entries] or [0]) + 1

    created, updated, copied, missing_files = [], [], 0, []
    for row in rows:
        if row["include"] != YES or row.get("folder"):
            continue
        rid = row.get("row_id") or "manual:%s" % sanitize(row.get("property") or "row")
        row["row_id"] = rid
        existing = by_row.get(rid)
        order = int(existing["order"]) if existing else next_order
        folder = existing["folder"] if existing else property_folder(
            order, sanitize(row.get("property") or "Property"), (row.get("postcode") or "").strip() or None)
        pdir = os.path.join(props_dir, folder)
        os.makedirs(os.path.join(pdir, "media", "images"), exist_ok=True)

        d = synth_derived(row)
        files = row.get("files") or manifest_files.get(rid) or []
        if isinstance(files, str):
            files = [x.strip() for x in re.split(r"[;\n]", files) if x.strip()]
        for ref in files:
            src = resolve_file(work, ref)
            if not src:
                missing_files.append((row.get("property"), ref))
                continue
            dest = os.path.join(pdir, "media", os.path.basename(src))
            if not os.path.exists(dest):
                shutil.copy2(src, dest)
                copied += 1
            d["documents"].append({"name": os.path.basename(src), "ext": os.path.splitext(src)[1].lstrip("."),
                                   "size": os.path.getsize(dest), "kind": "Brochure", "url": None})
        write_json(os.path.join(pdir, "_derived.json"), d)
        entry = {"order": order, "match_id": None, "folder": folder,
                 "name": row.get("property"), "postcode": (row.get("postcode") or "").strip() or None,
                 "for_sale": d["for_sale"], "to_let": d["to_let"], "group_position": None,
                 "origin": "master_list", "row_id": rid,
                 "source_type": row.get("source_type") or "Broker email",
                 "source": row.get("source") or ""}
        if existing:
            existing.update(entry)
            updated.append(folder)
        else:
            entries.append(entry)
            by_row[rid] = entry
            next_order += 1
            created.append(folder)
        row["folder"] = folder

    if created or updated:
        entries.sort(key=lambda e: int(e.get("order") or 0))
        idx["properties"] = entries
        idx["count"] = len(entries)
        write_json(idx_path, idx)
    return created, updated, copied, missing_files


# ------------------------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", required=True)
    ap.add_argument("--workbook", default=None, help="path to the edited master list (default: found in <work>)")
    ap.add_argument("--no-materialise", action="store_true",
                    help="do not create property folders for included email/upload-only options. "
                         "They will be recorded in master_list.json and built by nothing.")
    args = ap.parse_args()

    cfg = load_config(args.config)
    work = cfg["work_dir"]
    path = find_workbook(work, args.workbook)
    print("Reading %s" % path, flush=True)

    manifest = read_json(os.path.join(work, B.MANIFEST), {}) or {}
    man_rows = {r.get("row_id"): r for r in (manifest.get("rows") or [])}
    manifest_files = {rid: (r.get("files") or []) for rid, r in man_rows.items()}

    rows = read_rows(path)
    if not rows:
        raise SystemExit("%s has no data rows." % path)

    blank, deferred, unknown = [], [], []
    for r in rows:
        r["include"] = norm_include(r.get("include"))
        rid = (r.get("row_id") or "").strip()
        r["row_id"] = rid or None
        man = man_rows.get(rid) or {}
        r["folder"] = man.get("folder")
        r["source_type"] = r.get("source_type") or man.get("source_type")
        label = "row %d  %s" % (r["_excel_row"], r.get("property") or "(no name)")
        if r["include"] == "":
            blank.append(label)
        elif r["include"] in (YES, NO):
            continue
        elif str(r["include"]).strip().lower() in _DEFERRED:
            deferred.append("%s  -> %r" % (label, r["include"]))
        else:
            unknown.append("%s  -> %r" % (label, r["include"]))

    if blank or deferred or unknown:
        print("ERROR: the master list is not fully answered, so nothing has been changed.",
              file=sys.stderr)
        for lbl in blank:
            print("  blank Include?      %s" % lbl, file=sys.stderr)
        for lbl in deferred:
            print("  not an answer here  %s" % lbl, file=sys.stderr)
        for lbl in unknown:
            print("  unreadable value    %s" % lbl, file=sys.stderr)
        if deferred:
            print("", file=sys.stderr)
            print("  Every row is Yes or No. A deferred answer has to be resolved before the run "
                  "can", file=sys.stderr)
            print("  start in any case, so it cannot be carried in the sheet: it would only move "
                  "the same", file=sys.stderr)
            print("  decision to a point where a log reader takes it instead of the person who "
                  "owns the", file=sys.stderr)
            print("  deliverable. Put the row(s) above to the user.", file=sys.stderr)
        print("", file=sys.stderr)
        print("  Set every Include? to Yes or No and re-run. Nothing was written.", file=sys.stderr)
        sys.exit(2)

    seen = {r["row_id"] for r in rows if r.get("row_id")}
    deleted = [r for rid, r in man_rows.items() if rid and rid not in seen]

    created = updated = copied = 0
    missing_files, new_folders = [], []
    if not args.no_materialise:
        new_folders, upd, copied, missing_files = materialise(work, rows, manifest_files)
        created, updated = len(new_folders), len(upd)

    merge, kept_separate, not_merged = build_merge_map(
        rows, manifest.get("duplicate_groups") or {})

    included = [r for r in rows if r["include"] == YES]
    excluded = [r for r in rows if r["include"] == NO]
    no_folder = [r for r in included if not r.get("folder")]

    out = {
        "generated": datetime.datetime.now().isoformat(timespec="seconds"),
        "source_workbook": os.path.basename(path),
        "counts": {"rows": len(rows), "included": len(included), "excluded": len(excluded),
                   "materialised": created, "merged_siblings": len(merge)},
        "included_folders": [r["folder"] for r in included if r.get("folder")],
        "excluded_folders": [r["folder"] for r in excluded if r.get("folder")],
        # <excluded folder> -> the included folder it is merged into. build_dataset.py drops the
        # excluded record and hands its folder to the survivor as a _dedupe_folders sibling, so
        # its brochure still reaches the pipeline.
        "merge_map": {k: v["into"] for k, v in merge.items() if v.get("into")},
        "merge_detail": merge,
        # Groups where the user deliberately kept more than one row. Recorded so the automatic
        # postcode/size merge in common.dedupe_props cannot quietly undo that decision.
        "kept_separate": [{"group": g, "properties": p} for g, p in kept_separate],
        # Rejected rows inside a group too loosely matched to merge. Their documents do NOT
        # travel to the kept row; named here so the loss is on the record rather than silent.
        "excluded_not_merged": [{"group": g, "property": p, "kept": k} for g, p, k in not_merged],
        "run_notes": {r["folder"]: r.get("run_notes") for r in included
                      if r.get("folder") and (r.get("run_notes") or "").strip()},
        "rows": [{k: r.get(k) for k in
                  ("row_id", "rank", "include", "run_notes", "property", "source_type", "source",
                   "folder", "postcode", "duplicate_group", "duplicate_status", "_excel_row")}
                 for r in rows],
        "deleted_since_build": [{"row_id": r.get("row_id"), "property": r.get("property"),
                                 "folder": r.get("folder")} for r in deleted],
        "unresolved_files": [{"property": p, "file": f} for p, f in missing_files],
    }
    write_json(os.path.join(work, MASTER_JSON), out)

    notes_n = len(out["run_notes"])
    print("MASTER LIST READ -> %s" % os.path.join(work, MASTER_JSON), flush=True)
    print("  include=%d  exclude=%d  run notes=%d  merged siblings=%d"
          % (len(included), len(excluded), notes_n, len(merge)), flush=True)
    if created or updated:
        print("  materialised %d new propert(y/ies), updated %d, copied %d document(s):"
              % (created, updated, copied), flush=True)
        for f in new_folders:
            print("    + properties/%s" % f, flush=True)
    if no_folder:
        print("  WARNING: %d included row(s) have no property folder and will be built by NOTHING: %s"
              % (len(no_folder), "; ".join(r.get("property") or "?" for r in no_folder)), file=sys.stderr)
    if missing_files:
        print("  WARNING: %d file(s) named on the master list were not found on disk: %s"
              % (len(missing_files), "; ".join("%s -> %s" % (p, f) for p, f in missing_files[:5])),
              file=sys.stderr)
    if deleted:
        print("  WARNING: %d row(s) from the build are not in the workbook and are treated as "
              "excluded: %s" % (len(deleted), "; ".join(r.get("property") or "?" for r in deleted)),
              file=sys.stderr)
    for g, p, k in not_merged:
        print("  group %s: '%s' excluded and NOT merged into '%s' (postcode match only, areas "
              "differ) - any document it held is not used" % (g, p, k), file=sys.stderr)
    for g, props in kept_separate:
        print("  group %s kept as %d separate propert(y/ies) by the user: %s"
              % (g, len(props), "; ".join(p or "?" for p in props)), flush=True)
    if notes_n:
        print("  NEXT: read master_list.json run_notes before enrichment - they are instructions "
              "for you, not data.", flush=True)


if __name__ == "__main__":
    main()
